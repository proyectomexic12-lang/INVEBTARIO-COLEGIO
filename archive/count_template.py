import openpyxl
wb = openpyxl.load_workbook('PlantillaConsultaInventarioGeneral1.xlsx', data_only=True)
sheet = wb.active
rows = list(sheet.iter_rows(values_only=True))
valid_data_rows = 0
for r in rows:
    # Try row content check
    if any(str(c).strip() and str(c).upper() not in ["NONE", ""] for c in r):
        # But skip headers (approx row 6)
        if str(r[0]).upper() not in ["CÓDIGO", "NONE", "ID", "PLACA"]:
            valid_data_rows += 1
print(f"Total Rows with data in Template: {valid_data_rows}")
# Check if there are any sheets that say 'BAJAS'
print(f"Sheet names: {wb.sheetnames}")
