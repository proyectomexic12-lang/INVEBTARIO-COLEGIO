import openpyxl
from database import db
import os
from datetime import datetime

def clean_num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('$', '').replace(' ', '')
    if not s: return 0.0
    try:
        # Detect standard English format vs Spanish/Colombian format
        if ',' in s and '.' in s:
            if s.rfind(',') > s.rfind('.'):
                s = s.replace('.', '').replace(',', '.')
            else:
                s = s.replace(',', '')
        elif ',' in s:
            last_comma = s.rfind(',')
            if len(s) - last_comma - 1 == 3: # thousands
                s = s.replace(',', '')
            else: # decimal
                s = s.replace(',', '.')
        elif '.' in s:
            last_dot = s.rfind('.')
            if len(s) - last_dot - 1 == 3: # thousands
                s = s.replace('.', '')
        return float(s)
    except Exception:
        return 0.0

def import_real_data():
    possible_paths = [
        r"C:\Users\USUARIO\Downloads\inventario_guaimaral.xlsx",
        os.path.join(os.path.expanduser("~"), "Downloads", "inventario_guaimaral.xlsx"),
        os.path.join(os.path.expanduser("~"), "Desktop", "inventario_guaimaral.xlsx"),
        os.path.join(os.path.dirname(__file__), "inventario_guaimaral.xlsx"),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), "inventario_guaimaral.xlsx")
    ]
    template_path = None
    for path in possible_paths:
        if os.path.exists(path):
            template_path = path
            break
            
    if not template_path:
        print("Template not found at any of these paths:")
        for path in possible_paths:
            print(f"  - {path}")
        return


    print("Cargando libro de Excel...")
    wb = openpyxl.load_workbook(template_path, data_only=True)
    sheet = wb.active
    
    # 1. Clean existing data
    print("Limpiando tablas de activos, préstamos y dependencias en la base de datos...")
    db.cursor.execute("DELETE FROM assets")
    db.cursor.execute("DELETE FROM dependencies")
    db.cursor.execute("DELETE FROM loans")
    db.conn.commit()

    # Column name to index mapping based on inventario_guaimaral.xlsx header row
    headers = [str(c.value).strip().upper() if c.value else "" for c in sheet[1]]
    mapping = {name: i for i, name in enumerate(headers) if name}
    
    def gx(row, key, default=""):
        idx = mapping.get(key)
        if idx is not None and idx < len(row) and row[idx] is not None:
            return str(row[idx]).strip()
        return default

    # 2. Extract unique dependencies
    print("Identificando dependencias...")
    deps = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        dep = gx(row, "UBICACION")
        if dep:
            deps.add(dep.upper())
    
    print(f"Registrando {len(deps)} dependencias...")
    for d in sorted(deps):
        db.add_dependency(d, "AULA")

    # 3. Import Assets
    print("Importando activos...")
    assets_imported = 0
    batch = []
    now = datetime.now().strftime("%Y-%m-%d")

    for row in sheet.iter_rows(min_row=2, values_only=True):
        desc = gx(row, "DESCRIPCION")
        if not desc or desc.upper() in ["DESCRIPCIÓN", "ARTÍCULO", "NONE", "NULL", ""]:
            continue
            
        qty_val = int(clean_num(gx(row, "EXISTENCIAINICIAL")))
        qty = max(1, qty_val)
        val_uni = clean_num(gx(row, "VALOR"))
        val_tot = val_uni * qty
        depr = clean_num(gx(row, "DEPRECIACUMULADA"))
        val_cur = max(0.0, val_tot - depr)
        
        try: life = int(clean_num(gx(row, "VIDAUTIL")))
        except: life = 10

        serial_val = gx(row, "SERIAL", "N/A")
        if serial_val.upper() in ["NULL", "NONE", "—", "SIN", "S/N", ""]:
            serial_val = "N/A"

        code_val = gx(row, "CODIGO")
        if not code_val or code_val.upper() in ["NULL", "NONE", "—", ""]:
            code_val = f"NX-GUA-{assets_imported+1:04d}"

        st = "Bueno"
        act = "En Uso"
        is_disposed = 0

        data = (
            code_val,
            gx(row, "CODCONTABLE", "101"),      # acc_code
            "GENERAL",                          # rubro description
            desc,                               # description
            gx(row, "MARCA", "N/A"),            # brand
            "N/A",                              # model
            serial_val,                         # serial_number
            "N/A",                              # color
            "N/A",                              # dimensions
            "Migrado de Excel",                  # observations
            val_uni,                            # initial_unit_value
            qty,                                # quantity
            val_tot,                            # initial_total_value
            depr,                               # depreciation_rate
            val_cur,                            # current_value
            gx(row, "FECHAADQUISICION", now),   # entry_date
            gx(row, "FECHAADQUISICION", now),   # service_date
            life,                               # useful_life_remaining
            st,                                 # conservation_state
            "Propio",                           # origin
            act,                                # current_activity
            gx(row, "UBICACION", "SIN ASIGNAR").upper(), # dependency
            is_disposed,
            None,
            None,
            gx(row, "CODDEPRECIACION"),
            gx(row, "CODGASTO"),
            gx(row, "FUNCIONARIO"),
            gx(row, "IDENTIFICACION"),
            gx(row, "CODGRUPO"),
            gx(row, "CODSUBGRUPO"),
            gx(row, "TIPO"),
            gx(row, "SEDE", "Guaimaral")
        )
        batch.append(data)
        assets_imported += 1

    db.bulk_add_assets(batch)
    print(f"¡Importación completada! {len(deps)} dependencias y {assets_imported} activos agregados.")

if __name__ == "__main__":
    import_real_data()
