import customtkinter as ctk
import tkinter as tk
from tkinter import messagebox
from datetime import datetime
import re
from theme import theme
from scroll_helper import enable_smooth_scroll
class ModernInventory(ctk.CTkFrame):
    def __init__(self, master, controller, asset_repo, dep_repo):
        super().__init__(master, fg_color="transparent")
        self.controller = controller
        self.asset_repo = asset_repo
        self.dep_repo = dep_repo
        self.current_page = 1
        self.items_per_page = 15
        self.assets = []
        
        # --- Header ---
        self.header = ctk.CTkFrame(self, fg_color=theme.BG_CARD, corner_radius=20, border_width=1, border_color=theme.BORDER)
        self.header.pack(fill="x", padx=40, pady=(35, 10))
        h_inner = ctk.CTkFrame(self.header, fg_color="transparent")
        h_inner.pack(padx=30, pady=20, fill="x")
        
        ctk.CTkLabel(h_inner, text="ENTORNO DE GESTIÓN GLOBAL", font=ctk.CTkFont(size=10, weight="bold"), text_color=theme.PRIMARY_LIGHT).pack(anchor="w")
        ctk.CTkLabel(h_inner, text="Explorador Central de Activos", font=ctk.CTkFont(size=28, weight="bold"), text_color=theme.TEXT_MAIN).pack(anchor="w")
        
        # --- Advanced Filters (Task 16) ---
        self.setup_filters()
        
        # --- Pagination Control ---
        self.pag_bar = ctk.CTkFrame(self, fg_color="transparent")
        self.pag_bar.pack(fill="x", padx=40, pady=8)
        
        self.prev_btn = ctk.CTkButton(self.pag_bar, text="◀ ANTERIOR", width=120, height=36, corner_radius=10, 
                                     fg_color=theme.BG_CARD, hover_color=theme.PRIMARY_HOVER, border_width=1, border_color=theme.BORDER,
                                     text_color=theme.TEXT_MAIN, font=ctk.CTkFont(size=11, weight="bold"),
                                     command=lambda: self.change_page(-1))
        self.prev_btn.pack(side="left")
        
        self.page_label = ctk.CTkLabel(self.pag_bar, text="Página 1", font=ctk.CTkFont(size=13, weight="bold"), text_color=theme.TEXT_MUTED)
        self.page_label.pack(side="left", expand=True)
        
        self.next_btn = ctk.CTkButton(self.pag_bar, text="SIGUIENTE ▶", width=120, height=36, corner_radius=10, 
                                     fg_color=theme.BG_CARD, hover_color=theme.PRIMARY_HOVER, border_width=1, border_color=theme.BORDER,
                                     text_color=theme.TEXT_MAIN, font=ctk.CTkFont(size=11, weight="bold"),
                                     command=lambda: self.change_page(1))
        self.next_btn.pack(side="right")
        
        # --- Bulk Actions Frame ---
        self.bulk_bar = ctk.CTkFrame(self, fg_color=theme.BG_CARD, corner_radius=15, border_width=1, border_color=theme.BORDER, height=55)
        self.bulk_bar.pack(fill="x", padx=40, pady=5)
        self.bulk_bar.pack_propagate(False)
        
        self.bulk_lbl = ctk.CTkLabel(self.bulk_bar, text="Acciones en Lote (0 seleccionados)", font=ctk.CTkFont(size=12, weight="bold"), text_color=theme.TEXT_MUTED)
        self.bulk_lbl.pack(side="left", padx=20)
        
        self.bulk_transfer_btn = ctk.CTkButton(self.bulk_bar, text="🔄 Traslado Masivo", width=145, height=34, corner_radius=8, fg_color=theme.WARNING, hover_color="#d97706", state="disabled", font=ctk.CTkFont(size=11, weight="bold"), command=self.bulk_transfer)
        self.bulk_transfer_btn.pack(side="right", padx=10, pady=10)
        
        self.bulk_dispose_btn = ctk.CTkButton(self.bulk_bar, text="📉 Baja Masiva", width=120, height=34, corner_radius=8, fg_color=theme.DANGER, hover_color="#dc2626", state="disabled", font=ctk.CTkFont(size=11, weight="bold"), command=self.bulk_dispose)
        self.bulk_dispose_btn.pack(side="right", padx=6, pady=10)
        
        self.bulk_delete_btn = ctk.CTkButton(self.bulk_bar, text="🗑️ Eliminar Lote", width=125, height=34, corner_radius=8, fg_color="#991b1b", hover_color="#7f1d1d", state="disabled", font=ctk.CTkFont(size=11, weight="bold"), command=self.bulk_delete)
        self.bulk_delete_btn.pack(side="right", padx=6, pady=10)

        # --- Table Frame (Scrollable) ---
        self.table_wrapper = ctk.CTkFrame(self, fg_color=theme.BG_CARD, corner_radius=18, border_width=1, border_color=theme.BORDER)
        self.table_wrapper.pack(fill="both", expand=True, padx=40, pady=10)
        
        self.canvas = tk.Canvas(self.table_wrapper, bg=theme.BG_CARD, highlightthickness=0)
        self.h_sb = ctk.CTkScrollbar(self.table_wrapper, orientation="horizontal", command=self.canvas.xview)
        self.v_sb = ctk.CTkScrollbar(self.table_wrapper, orientation="vertical", command=self.canvas.yview)
        self.canvas.configure(xscrollcommand=self.h_sb.set, yscrollcommand=self.v_sb.set)
        
        self.v_sb.pack(side="right", fill="y")
        self.h_sb.pack(side="bottom", fill="x")
        self.canvas.pack(side="left", fill="both", expand=True)
        
        self.table_content = ctk.CTkFrame(self.canvas, fg_color=theme.BG_CARD)
        self.canvas_win = self.canvas.create_window((0, 0), window=self.table_content, anchor="nw")
        
        def _update_inv_scrollregion(event=None):
            bbox = self.canvas.bbox("all")
            if bbox:
                w = max(bbox[2], self.canvas.winfo_width())
                h = max(bbox[3], self.canvas.winfo_height())
                self.canvas.configure(scrollregion=(0, 0, w, h))

        self.table_content.bind("<Configure>", _update_inv_scrollregion)
        self.canvas.bind("<Configure>", _update_inv_scrollregion)
        
        # Universal Smooth Mousewheel Scrolling
        enable_smooth_scroll(self.canvas, self.table_wrapper)
        
        self.total_count = 0
        self.is_search_mode = False
        self.refresh_grid()

    def refresh_data(self):
        self.refresh_grid()

    def inject_results(self, results):
        self.is_search_mode = True
        self.search_results_all = results
        self.total_count = len(results)
        self.current_page = 1
        self._render_search_page()

    def _render_search_page(self):
        start = (self.current_page - 1) * self.items_per_page
        end = start + self.items_per_page
        page_results = self.search_results_all[start:end]
        sum_ini = 0.0
        sum_depr = 0.0
        for a in self.search_results_all:
            try:
                v = float(str(a.get('VALOR', 0)).replace('$', '').replace(',', '').strip() or 0)
                q = max(1, int(str(a.get('EXISTENCIAINICIAL', 1)).replace(',', '').strip() or 1))
                sum_ini += (v * q)
            except: pass
            try:
                d = float(str(a.get('DEPRECIACUMULADA', 0)).replace('$', '').replace(',', '').strip() or 0)
                sum_depr += d
            except: pass
        self._render_grid_callback((page_results, self.total_count, sum_ini, sum_depr))

    def change_page(self, delta):
        self.current_page += delta
        if self.is_search_mode:
            self._render_search_page()
        else:
            self.refresh_grid()

    def refresh_grid(self):
        if self.is_search_mode: return
        self.update_filter_comboboxes()
        offset = (self.current_page - 1) * self.items_per_page
        
        # Read filters on MAIN thread (prevents Tkinter threading RuntimeError)
        dep = self.dep_filter.get() if hasattr(self, "dep_filter") else "Todas"
        state = self.state_filter.get() if hasattr(self, "state_filter") else "Todos"
        
        min_v = None
        if hasattr(self, "min_val_entry"):
            min_raw = self.min_val_entry.get().strip()
            if min_raw:
                try: min_v = float(min_raw.replace('$','').replace(',','').replace(' ',''))
                except: pass
                
        max_v = None
        if hasattr(self, "max_val_entry"):
            max_raw = self.max_val_entry.get().strip()
            if max_raw:
                try: max_v = float(max_raw.replace('$','').replace(',','').replace(' ',''))
                except: pass
                
        sede = getattr(self.controller, "global_sede", "Todas")
        self.controller.run_in_thread(
            self._fetch_inventory_data,
            self._render_grid_callback,
            offset, dep, state, min_v, max_v, sede
        )

    def _fetch_inventory_data(self, offset, dep, state, min_v, max_v, sede):
        assets, total, sum_ini, sum_depr = self.asset_repo.get_filtered_assets(
            limit=self.items_per_page,
            offset=offset,
            dependency=dep,
            state=state,
            min_val=min_v,
            max_val=max_v,
            sede=sede
        )
        return assets, total, sum_ini, sum_depr

    def setup_filters(self):
        self.filter_bar = ctk.CTkFrame(self, fg_color=theme.BG_CARD, corner_radius=18, border_width=1, border_color=theme.BORDER)
        self.filter_bar.pack(fill="x", padx=40, pady=5)
        
        inner_fb = ctk.CTkFrame(self.filter_bar, fg_color="transparent")
        inner_fb.pack(padx=20, pady=12, fill="x")
        
        inner_fb.columnconfigure(0, weight=3) # Dependency
        inner_fb.columnconfigure(1, weight=2) # State
        inner_fb.columnconfigure(2, weight=1) # Actions
        
        # 1. Dependency
        dep_f = ctk.CTkFrame(inner_fb, fg_color="transparent")
        dep_f.grid(row=0, column=0, padx=8, sticky="ew")
        ctk.CTkLabel(dep_f, text="DEPENDENCIA", font=ctk.CTkFont(size=9, weight="bold"), text_color=theme.TEXT_FAINT).pack(anchor="w", pady=(0, 4))
        
        db_deps = self.dep_repo.get_dependencies()
        self.dep_filter = ctk.CTkComboBox(dep_f, values=["Todas"] + db_deps, height=36, corner_radius=8, 
                                           border_width=1, border_color=theme.BORDER, fg_color=theme.BG_INPUT,
                                           command=lambda v: self.apply_filters())
        self.dep_filter.pack(fill="x")
        self.dep_filter.set("Todas")
        
        # 2. State
        state_f = ctk.CTkFrame(inner_fb, fg_color="transparent")
        state_f.grid(row=0, column=1, padx=8, sticky="ew")
        ctk.CTkLabel(state_f, text="ESTADO CONSERVACIÓN", font=ctk.CTkFont(size=9, weight="bold"), text_color=theme.TEXT_FAINT).pack(anchor="w", pady=(0, 4))
        
        self.state_filter = ctk.CTkComboBox(state_f, values=["Todos", "Excelente", "Bueno", "Regular", "Malo"], height=36, corner_radius=8,
                                             border_width=1, border_color=theme.BORDER, fg_color=theme.BG_INPUT,
                                             command=lambda v: self.apply_filters())
        self.state_filter.pack(fill="x")
        self.state_filter.set("Todos")
        
        # 3. Buttons
        btn_f = ctk.CTkFrame(inner_fb, fg_color="transparent")
        btn_f.grid(row=0, column=2, padx=8, sticky="ew")
        ctk.CTkLabel(btn_f, text="ACCIONES", font=ctk.CTkFont(size=9, weight="bold"), text_color=theme.TEXT_FAINT).pack(anchor="w", pady=(0, 4))
        
        sub_btn_f = ctk.CTkFrame(btn_f, fg_color="transparent")
        sub_btn_f.pack(fill="x")
        
        ctk.CTkButton(sub_btn_f, text="Filtrar", height=36, fg_color=theme.PRIMARY, hover_color=theme.PRIMARY_HOVER, corner_radius=8, font=ctk.CTkFont(size=11, weight="bold"), command=self.apply_filters).pack(side="left", fill="x", expand=True, padx=(0, 4))
        ctk.CTkButton(sub_btn_f, text="Limpiar", height=36, fg_color="transparent", border_width=1, border_color=theme.BORDER, text_color=theme.TEXT_MUTED, hover_color=theme.BG_SIDEBAR, corner_radius=8, font=ctk.CTkFont(size=11), command=self.clear_filters).pack(side="left", fill="x", expand=True, padx=(4, 0))

        # Global Summary Bar for Filtered Assets
        self.summary_bar = ctk.CTkFrame(self, fg_color=theme.BG_SIDEBAR, corner_radius=12, border_width=1, border_color=theme.BORDER)
        self.summary_bar.pack(fill="x", padx=40, pady=(0, 10))
        
        self.btn_export = ctk.CTkButton(self.summary_bar, text="📊 Exportar Excel", height=32, corner_radius=8,
                                        fg_color="#107C41", hover_color="#185C37", font=ctk.CTkFont(size=12, weight="bold"),
                                        command=self.export_to_excel)
        self.btn_export.pack(side="right", padx=15, pady=10)
        
        ctk.CTkLabel(self.summary_bar, text="BALANCE CONTABLE:", font=ctk.CTkFont(size=10, weight="bold"), text_color=theme.PRIMARY_LIGHT).pack(side="left", padx=(20, 10), pady=10)
        
        self.lbl_sum_initial = ctk.CTkLabel(self.summary_bar, text="VALOR EN ACTIVO: $0", font=ctk.CTkFont(size=12, weight="bold"), text_color=theme.SUCCESS)
        self.lbl_sum_initial.pack(side="left", padx=15, pady=10)
        
        self.lbl_sum_depr = ctk.CTkLabel(self.summary_bar, text="DEPRECIACIÓN ACUMULADA: $0", font=ctk.CTkFont(size=12, weight="bold"), text_color=theme.DANGER)
        self.lbl_sum_depr.pack(side="left", padx=15, pady=10)

        self.lbl_sum_real = ctk.CTkLabel(self.summary_bar, text="VALOR REAL DEL ACTIVO: $0", font=ctk.CTkFont(size=12, weight="bold"), text_color=theme.WARNING)
        self.lbl_sum_real.pack(side="left", padx=15, pady=10)

    def export_to_excel(self):
        import tkinter.filedialog as fd
        
        dep = self.dep_filter.get() if hasattr(self, "dep_filter") else "Todas"
        sede = getattr(self.controller, 'global_sede', 'Todas')
        
        if dep and dep != "Todas":
            filename = f"Inventario_{sede}_{dep.replace(' ', '_')}.xlsx"
        else:
            filename = f"Inventario_{sede}.xlsx"
            
        filepath = fd.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            title="Guardar Reporte Excel",
            initialfile=filename
        )
        if not filepath: return
        self.controller.run_in_thread(self._run_export_excel, self._on_export_done, filepath)

    def _on_export_done(self, success):
        if success:
            messagebox.showinfo("Exportación Exitosa", "El reporte Excel se ha guardado correctamente.")
        else:
            messagebox.showerror("Error", "Ocurrió un problema al generar el Excel. Asegúrese de que el archivo no esté abierto.")

    def _run_export_excel(self, filepath):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            # Fetch all filtered assets without pagination
            dep = self.dep_filter.get() if hasattr(self, "dep_filter") else "Todas"
            state = self.state_filter.get() if hasattr(self, "state_filter") else "Todos"
            
            min_v = None
            if hasattr(self, "min_val_entry"):
                try: min_v = float(self.min_val_entry.get().replace('$','').replace(',','').strip())
                except: pass
            max_v = None
            if hasattr(self, "max_val_entry"):
                try: max_v = float(self.max_val_entry.get().replace('$','').replace(',','').strip())
                except: pass
                
            sede = getattr(self.controller, "global_sede", "Todas")
            
            assets, _, _, _ = self.asset_repo.get_filtered_assets(
                limit=999999, offset=0, dependency=dep, state=state, min_val=min_v, max_val=max_v, sede=sede
            )
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Inventario Patrimonial"
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            headers = [
                "Código", "Descripción", "Fecha Adquisición", "Marca", "Serie", "Sede", "Ubicación (Salón)",
                "Cód Contable", "Vida Útil", "Cód Depreciación", "Gasto",
                "Funcionario", "Identificación", "Grupo", "Subgrupo", "Tipo", "Estado",
                "Cantidad", "Valor Unitario", "Valor Total", "Depreciación Acum.", "Valor en Libros"
            ]
            
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            
            for col_num, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                cell.border = thin_border
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 18
                
            # Group by dependency for summary
            dep_summary = {}
            
            def to_title(val):
                if not val: return ""
                return str(val).strip().title()
                
            for row_num, a in enumerate(assets, 2):
                qty = int(a.get('EXISTENCIAINICIAL') or 1)
                val_uni = float(a.get('VALOR') or 0.0)
                val_tot = val_uni * qty
                depr = float(a.get('DEPRECIACUMULADA') or 0.0)
                val_libros = max(0.0, val_tot - depr)
                
                ubicacion = to_title(a.get('UBICACION') or 'SIN ASIGNAR')
                
                if ubicacion not in dep_summary:
                    dep_summary[ubicacion] = {'qty': 0, 'val_tot': 0.0, 'depr': 0.0, 'val_libros': 0.0}
                
                dep_summary[ubicacion]['qty'] += qty
                dep_summary[ubicacion]['val_tot'] += val_tot
                dep_summary[ubicacion]['depr'] += depr
                dep_summary[ubicacion]['val_libros'] += val_libros
                
                row_data = [
                    a.get('CODIGO'), to_title(a.get('DESCRIPCION')), a.get('FECHAADQUISICION'), to_title(a.get('MARCA')), a.get('SERIAL'),
                    to_title(a.get('SEDE')), ubicacion, a.get('CODCONTABLE'), a.get('VIDAUTIL'),
                    a.get('CODDEPRECIACION'), a.get('CODGASTO'), to_title(a.get('FUNCIONARIO')),
                    a.get('IDENTIFICACION'), to_title(a.get('CODGRUPO')), to_title(a.get('CODSUBGRUPO')),
                    to_title(a.get('TIPO')), to_title(a.get('conservation_state')),
                    qty, val_uni, val_tot, depr, val_libros
                ]
                
                for col_num, val in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num, value=val)
                    cell.border = thin_border
                    # Align numbers to right, text to left
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")
                    if col_num in (19, 20, 21, 22):
                        cell.number_format = '"$"#,##0.00'
                    
            # -------------------------------------------------------------
            # SUMMARY BY DEPENDENCY (AT THE END)
            # -------------------------------------------------------------
            start_row = len(assets) + 4
            
            ws.cell(row=start_row, column=1, value="RESUMEN CONTABLE POR SALÓN / DEPENDENCIA").font = Font(bold=True, size=14)
            ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=5)
            
            start_row += 2
            sum_headers = ["Salón / Dependencia", "Total Activos", "Inversión Total", "Depreciación Acum.", "Valor en Libros Total"]
            for col_num, h in enumerate(sum_headers, 1):
                cell = ws.cell(row=start_row, column=col_num, value=h)
                cell.fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
                cell.font = header_font
                cell.border = thin_border
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20
                
            total_activos = 0
            total_inv = 0.0
            total_dep = 0.0
            total_libros = 0.0
                
            for dep_name in sorted(dep_summary.keys()):
                start_row += 1
                data = dep_summary[dep_name]
                
                total_activos += data['qty']
                total_inv += data['val_tot']
                total_dep += data['depr']
                total_libros += data['val_libros']
                
                c1 = ws.cell(row=start_row, column=1, value=dep_name)
                c2 = ws.cell(row=start_row, column=2, value=data['qty'])
                c3 = ws.cell(row=start_row, column=3, value=data['val_tot'])
                c4 = ws.cell(row=start_row, column=4, value=data['depr'])
                c5 = ws.cell(row=start_row, column=5, value=data['val_libros'])
                
                for c in (c1, c2, c3, c4, c5):
                    c.border = thin_border
                c3.number_format = '"$"#,##0.00'
                c4.number_format = '"$"#,##0.00'
                c5.number_format = '"$"#,##0.00'
                
            # GRAND TOTAL ROW
            start_row += 1
            c1 = ws.cell(row=start_row, column=1, value="TOTAL GENERAL")
            c2 = ws.cell(row=start_row, column=2, value=total_activos)
            c3 = ws.cell(row=start_row, column=3, value=total_inv)
            c4 = ws.cell(row=start_row, column=4, value=total_dep)
            c5 = ws.cell(row=start_row, column=5, value=total_libros)
            
            total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            for c in (c1, c2, c3, c4, c5):
                c.border = thin_border
                c.fill = total_fill
                c.font = Font(bold=True)
            c3.number_format = '"$"#,##0.00'
            c4.number_format = '"$"#,##0.00'
            c5.number_format = '"$"#,##0.00'
            
            wb.save(filepath)
            return True
        except Exception as e:
            import traceback
            traceback.print_exc()
            print(f"Excel Export Error: {e}")
            return False

    def apply_filters(self):
        self.is_search_mode = False
        self.current_page = 1
        self.refresh_grid()

    def clear_filters(self):
        self.dep_filter.set("Todas")
        self.state_filter.set("Todos")
        self.is_search_mode = False
        self.current_page = 1
        self.refresh_grid()

    def update_filter_comboboxes(self):
        if hasattr(self, "dep_filter"):
            sede = getattr(self.controller, "global_sede", "Todas")
            db_deps = self.dep_repo.get_dependencies(sede=sede)
            self.dep_filter.configure(values=["Todas"] + db_deps)

    def _render_grid_callback(self, data):
        self.assets, self.total_count, sum_ini, sum_depr = data
        
        if hasattr(self, "lbl_sum_initial"):
            val_real = max(0.0, sum_ini - sum_depr)
            self.lbl_sum_initial.configure(text=f"VALOR EN ACTIVO: ${sum_ini:,.0f}")
            self.lbl_sum_depr.configure(text=f"DEPRECIACIÓN ACUMULADA: ${sum_depr:,.0f}")
            if hasattr(self, "lbl_sum_real"):
                self.lbl_sum_real.configure(text=f"VALOR REAL DEL ACTIVO: ${val_real:,.0f}")
            
        for w in self.table_content.winfo_children(): w.destroy()
        
        self.checked_variables = {}
        max_p = max(1, (self.total_count - 1) // self.items_per_page + 1)
        self.page_label.configure(text=f"Página {self.current_page} de {max_p} ({self.total_count} Activos)")
        self.prev_btn.configure(state="normal" if self.current_page > 1 else "disabled")
        self.next_btn.configure(state="normal" if self.current_page < max_p else "disabled")
        
        headers = [
            ("", 40),
            ("DESCRIPCION", 380),
            ("VALOR", 120),
            ("FECHAADQUISICION", 130),
            ("MARCA", 130),
            ("SERIAL", 140),
            ("UBICACION", 160),
            ("CODCONTABLE", 120),
            ("VIDAUTIL", 90),
            ("CODDEPRECIACION", 130),
            ("DEPRECIACUMULADA", 140),
            ("CODGASTO", 120),
            ("FUNCIONARIO", 180),
            ("IDENTIFICACION", 130),
            ("CODGRUPO", 100),
            ("CODSUBGRUPO", 110),
            ("TIPO", 80),
            ("EXISTENCIAINICIAL", 120),
            ("SEDE", 150),
            ("ACCIONES", 310)
        ]
        
        h_row = ctk.CTkFrame(self.table_content, fg_color=theme.BG_SIDEBAR, corner_radius=0, height=45)
        h_row.pack(fill="x")
        
        self.select_all_var = tk.BooleanVar(value=False)
        self.select_all_cb = ctk.CTkCheckBox(h_row, text="", variable=self.select_all_var, width=40, height=45, command=self.toggle_select_all)
        self.select_all_cb.grid(row=0, column=0, padx=5)
        
        for j, (txt, w) in enumerate(headers[1:], 1):
            ctk.CTkLabel(h_row, text=txt, width=w, height=45, font=ctk.CTkFont(size=9, weight="bold"), text_color=theme.TEXT_MUTED, fg_color=theme.BG_SIDEBAR).grid(row=0, column=j)

        for i, a_row in enumerate(self.assets):
            a = dict(a_row)
            
            # Alerta de Vida Útil Crítica (Task 11)
            is_critical = False
            try:
                useful_life = int(a.get('VIDAUTIL') or 10)
                if useful_life <= 1:
                    is_critical = True
            except:
                pass
                
            bg = "#220e14" if is_critical else (theme.BG_CARD if i % 2 == 0 else theme.BG_APP)
            r_row = ctk.CTkFrame(self.table_content, fg_color=bg, corner_radius=0, height=50)
            r_row.pack(fill="x")
            
            cb_var = tk.BooleanVar(value=False)
            self.checked_variables[a['id']] = cb_var
            cb = ctk.CTkCheckBox(r_row, text="", variable=cb_var, width=40, height=50, command=self.update_bulk_bar)
            cb.grid(row=0, column=0, padx=5)
            
            # Formatear VALOR
            try:
                v_num = float(str(a.get('VALOR', 0)).replace('$', '').replace(',', '').strip() or 0)
                v_str = f"${v_num:,.0f}"
            except:
                v_str = str(a.get('VALOR') or '$0')

            # Formatear DEPRECIACUMULADA
            try:
                d_num = float(str(a.get('DEPRECIACUMULADA', 0)).replace('$', '').replace(',', '').strip() or 0)
                d_str = f"${d_num:,.0f}"
            except:
                d_str = str(a.get('DEPRECIACUMULADA') or '$0')

            desc_text = str(a.get('DESCRIPCION') or '—')
            if a.get('is_loaned'):
                desc_text = f"📦 [PRESTADO] {desc_text}"

            row_data = [
                (desc_text, 380),
                (v_str, 120),
                (a.get('FECHAADQUISICION') or '—', 130),
                (a.get('MARCA') or '—', 130),
                (a.get('SERIAL') or '—', 140),
                (a.get('UBICACION') or '—', 160),
                (a.get('CODCONTABLE') or '—', 120),
                (a.get('VIDAUTIL') if a.get('VIDAUTIL') is not None else '—', 90),
                (a.get('CODDEPRECIACION') or '—', 130),
                (d_str, 140),
                (a.get('CODGASTO') or '—', 120),
                (a.get('FUNCIONARIO') or '—', 180),
                (a.get('IDENTIFICACION') or '—', 130),
                (a.get('CODGRUPO') or '—', 100),
                (a.get('CODSUBGRUPO') or '—', 110),
                (a.get('TIPO') or '—', 80),
                (a.get('EXISTENCIAINICIAL') if a.get('EXISTENCIAINICIAL') is not None else '1', 120),
                (a.get('SEDE') or '—', 150)
            ]
            
            for j, (val, w) in enumerate(row_data, 1):
                ctk.CTkLabel(r_row, text=str(val)[:80], width=w, height=50, anchor="w", padx=10, font=ctk.CTkFont(size=11)).grid(row=0, column=j)

            # Columna de ACCIONES (Exactamente en la columna 19)
            actions_f = ctk.CTkFrame(r_row, fg_color="transparent", width=310, height=50)
            actions_f.grid(row=0, column=len(row_data)+1, padx=10)
            actions_f.pack_propagate(False)
            
            btn_edit = ctk.CTkButton(actions_f, text="Editar", width=65, height=28, corner_radius=6,
                                     fg_color=theme.PRIMARY, hover_color=theme.PRIMARY_HOVER, font=ctk.CTkFont(size=11, weight="bold"),
                                     command=lambda asset=a: self.open_edit(asset))
            btn_edit.pack(side="left", padx=2, pady=11)
            
            btn_trans = ctk.CTkButton(actions_f, text="Traslado", width=65, height=28, corner_radius=6,
                                      fg_color=theme.WARNING, hover_color="#d97706", font=ctk.CTkFont(size=11, weight="bold"),
                                      command=lambda asset=a: self.open_transfer(asset))
            btn_trans.pack(side="left", padx=2, pady=11)
            
            btn_baja = ctk.CTkButton(actions_f, text="Baja", width=65, height=28, corner_radius=6,
                                     fg_color=theme.DANGER, hover_color="#dc2626", font=ctk.CTkFont(size=11, weight="bold"),
                                     command=lambda asset=a: self.open_dispose(asset))
            btn_baja.pack(side="left", padx=2, pady=11)

            btn_del = ctk.CTkButton(actions_f, text="Eliminar", width=68, height=28, corner_radius=6,
                                    fg_color="#991b1b", hover_color="#7f1d1d", font=ctk.CTkFont(size=11, weight="bold"),
                                    command=lambda asset=a: self.direct_delete(asset))
            btn_del.pack(side="left", padx=2, pady=11)

        self.update_bulk_bar()

        self.table_content.update_idletasks()
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def open_edit(self, asset):
        def on_success():
            self.refresh_grid()
            self.controller.invalidate_all_pages()
        EditAssetDialog(self, asset, on_success)

    def open_transfer(self, asset):
        def on_success():
            self.refresh_grid()
            self.controller.invalidate_all_pages()
        TransferAssetDialog(self, asset, on_success)

    def open_dispose(self, asset):
        def on_success():
            self.refresh_grid()
            self.controller.invalidate_all_pages()
        DisposeAssetDialog(self, asset, on_success)

    def direct_delete(self, asset):
        desc = asset.get('DESCRIPCION') or asset.get('description') or 'este activo'
        if messagebox.askyesno("Confirmar Eliminación", f"¿Está seguro de que desea eliminar definitivamente '{desc}'?\n\nEsta acción no se puede deshacer y actualizará los balances contables."):
            asset_id = asset.get('id')
            if asset_id and self.asset_repo.delete_asset(asset_id):
                messagebox.showinfo("Activo Eliminado", f"El activo '{desc}' fue eliminado correctamente.")
                self.refresh_grid()
                self.controller.invalidate_all_pages()
                self.update_bulk_bar()
            else:
                messagebox.showerror("Error", "No se pudo eliminar el activo.")

    def toggle_select_all(self):
        state = self.select_all_var.get()
        for var in self.checked_variables.values():
            var.set(state)
        self.update_bulk_bar()

    def update_bulk_bar(self):
        selected_ids = [aid for aid, var in self.checked_variables.items() if var.get()]
        cnt = len(selected_ids)
        self.bulk_lbl.configure(text=f"Acciones en Lote ({cnt} seleccionados)", text_color="#3b82f6" if cnt > 0 else "#64748b")
        if cnt > 0:
            self.bulk_transfer_btn.configure(state="normal")
            self.bulk_dispose_btn.configure(state="normal")
            self.bulk_delete_btn.configure(state="normal")
        else:
            self.bulk_transfer_btn.configure(state="disabled")
            self.bulk_dispose_btn.configure(state="disabled")
            self.bulk_delete_btn.configure(state="disabled")

    def bulk_transfer(self):
        selected_ids = [aid for aid, var in self.checked_variables.items() if var.get()]
        if not selected_ids: return
        def on_success():
            self.refresh_grid()
            self.controller.invalidate_all_pages()
            self.update_bulk_bar()
        BulkTransferDialog(self, selected_ids, on_success)

    def bulk_dispose(self):
        selected_ids = [aid for aid, var in self.checked_variables.items() if var.get()]
        if not selected_ids: return
        def on_success():
            self.refresh_grid()
            self.controller.invalidate_all_pages()
            self.update_bulk_bar()
        BulkDisposeDialog(self, selected_ids, on_success)

    def bulk_delete(self):
        selected_ids = [aid for aid, var in self.checked_variables.items() if var.get()]
        if not selected_ids: return
        from tkinter import messagebox
        if messagebox.askyesno(
            "Confirmar Eliminación Masiva", 
            f"¿Está seguro de que desea eliminar definitivamente los {len(selected_ids)} activos seleccionados del inventario?\n\nEsta acción no se puede deshacer y actualizará los balances de inmediato."
        ):
            if self.asset_repo.bulk_delete_assets(selected_ids):
                messagebox.showinfo("Éxito", f"Se eliminaron {len(selected_ids)} activos correctamente.")
                self.refresh_grid()
                self.controller.invalidate_all_pages()
                self.update_bulk_bar()
            else:
                messagebox.showerror("Error", "No se pudieron eliminar los activos seleccionados.")


class EditAssetDialog(ctk.CTkToplevel):
    def __init__(self, parent, asset_dict, on_success_callback):
        super().__init__(parent)
        self.parent = parent
        self.asset = asset_dict
        self.on_success = on_success_callback
        
        self.title(f"EDITAR ACTIVO: {self.asset.get('code')}")
        self.configure(fg_color="#0a0f1d")
        
        self.grab_set()
        self.focus()
        self.transient(parent)
        
        # Center dialog responsively
        self.update_idletasks()
        try:
            width = 850
            screen_h = self.winfo_screenheight()
            height = min(740, max(500, screen_h - 80))
            x = max(0, (self.winfo_screenwidth() - width) // 2)
            y = max(0, (screen_h - height) // 2)
            self.geometry(f"{width}x{height}+{x}+{y}")
        except Exception:
            self.geometry("850x700")
        
        # Scrollable container for the form
        self.scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.scroll.pack(fill="both", expand=True, padx=20, pady=20)
        
        # Title
        ctk.CTkLabel(self.scroll, text="FICHA TÉCNICA - EDICIÓN RÁPIDA", font=ctk.CTkFont(size=11, weight="bold"), text_color="#3b82f6").pack(anchor="w", padx=15)
        title_desc = self.asset.get('DESCRIPCION') or self.asset.get('description') or 'ACTIVO'
        ctk.CTkLabel(self.scroll, text=f"{title_desc}", font=ctk.CTkFont(size=22, weight="bold"), text_color="#f8fafc").pack(anchor="w", padx=15, pady=(0, 15))
        
        self.entries = {}
        
        # Section I: Identificación y Ubicación
        sec1 = self.create_section("I. IDENTIFICACIÓN Y UBICACIÓN", "🏢")
        self.create_field(sec1, "Descripción", "DESCRIPCION", self.asset.get("DESCRIPCION") or self.asset.get("description"), row=0, col=0, colspan=2)
        self.create_field(sec1, "Ubicación (Salón/Área)", "UBICACION", self.asset.get("UBICACION") or self.asset.get("dependency"), row=1, col=0)
        self.create_field(sec1, "Sede Educativa", "SEDE", self.asset.get("SEDE") or self.asset.get("sede"), row=1, col=1, values=self.parent.asset_repo.get_all_sedes())
        
        # Section II: Especificaciones
        sec2 = self.create_section("II. ESPECIFICACIONES", "🔧")
        self.create_field(sec2, "Marca", "MARCA", self.asset.get("MARCA") or self.asset.get("brand"), row=0, col=0)
        self.create_field(sec2, "Serial No.", "SERIAL", self.asset.get("SERIAL") or self.asset.get("serial_number"), row=0, col=1)
        self.create_field(sec2, "Tipo", "TIPO", self.asset.get("TIPO") or self.asset.get("tipo"), row=1, col=0)
        self.create_field(sec2, "Existencia Inicial (Cantidad)", "EXISTENCIAINICIAL", str(self.asset.get("EXISTENCIAINICIAL") if self.asset.get("EXISTENCIAINICIAL") is not None else self.asset.get("quantity", 1)), row=1, col=1)
        
        # Section III: Valorización y Ciclo de Vida
        sec3 = self.create_section("III. VALORIZACIÓN Y CONTABILIDAD", "💰")
        self.create_field(sec3, "Valor (Costo Adquisición)", "VALOR", str(self.asset.get("VALOR") if self.asset.get("VALOR") is not None else self.asset.get("initial_unit_value", 0)), row=0, col=0)
        self.create_field(sec3, "Fecha Adquisición", "FECHAADQUISICION", self.asset.get("FECHAADQUISICION") or self.asset.get("entry_date"), row=0, col=1)
        self.create_field(sec3, "Depreciación Acumulada", "DEPRECIACUMULADA", str(self.asset.get("DEPRECIACUMULADA") if self.asset.get("DEPRECIACUMULADA") is not None else self.asset.get("depreciation_rate", 0)), row=1, col=0)
        self.create_field(sec3, "Vida Útil (Años)", "VIDAUTIL", str(self.asset.get("VIDAUTIL") if self.asset.get("VIDAUTIL") is not None else self.asset.get("useful_life_remaining", 10)), row=1, col=1)
        
        # Section IV: Códigos Contables y Custodio
        sec4 = self.create_section("IV. CÓDIGOS CONTABLES Y CUSTODIO", "📋")
        self.create_field(sec4, "Código Contable", "CODCONTABLE", self.asset.get("CODCONTABLE") or self.asset.get("accounting_rubric_code"), row=0, col=0)
        self.create_field(sec4, "Código Depreciación", "CODDEPRECIACION", self.asset.get("CODDEPRECIACION") or self.asset.get("cod_depreciacion"), row=0, col=1)
        self.create_field(sec4, "Código Gasto", "CODGASTO", self.asset.get("CODGASTO") or self.asset.get("cod_gasto"), row=1, col=0)
        self.create_field(sec4, "Funcionario Responsable", "FUNCIONARIO", self.asset.get("FUNCIONARIO") or self.asset.get("funcionario"), row=1, col=1)
        self.create_field(sec4, "Identificación (CC/NIT)", "IDENTIFICACION", self.asset.get("IDENTIFICACION") or self.asset.get("identificacion"), row=2, col=0)
        self.create_field(sec4, "Código Grupo", "CODGRUPO", self.asset.get("CODGRUPO") or self.asset.get("cod_grupo"), row=2, col=1)
        self.create_field(sec4, "Código Subgrupo", "CODSUBGRUPO", self.asset.get("CODSUBGRUPO") or self.asset.get("cod_subgrupo"), row=3, col=0)
        
        # Footer Action Bar
        footer = ctk.CTkFrame(self.scroll, fg_color="transparent")
        footer.pack(fill="x", pady=25)
        
        ctk.CTkButton(footer, text="💾 GUARDAR CAMBIOS", font=ctk.CTkFont(size=14, weight="bold"), height=50, fg_color="#10b981", hover_color="#059669", corner_radius=12, command=self.save).pack(side="right", padx=6, fill="x", expand=True)
        ctk.CTkButton(footer, text="🗑️ Eliminar Activo", font=ctk.CTkFont(size=13, weight="bold"), height=50, fg_color="#dc2626", hover_color="#b91c1c", corner_radius=12, command=self.delete_current_asset).pack(side="right", padx=6)
        ctk.CTkButton(footer, text="Cancelar", font=ctk.CTkFont(size=14), height=50, fg_color="#1e293b", hover_color="#334155", corner_radius=12, command=self.destroy).pack(side="left", padx=6)

        # Real-time bindings for dynamic depreciation calculation
        for key in ["FECHAADQUISICION", "VALOR", "VIDAUTIL", "EXISTENCIAINICIAL", "DESCRIPCION", "UBICACION"]:
            if key in self.entries:
                self.entries[key].bind("<KeyRelease>", self._live_calc_depreciation)

    def _live_calc_depreciation(self, event=None):
        try:
            def v(k): return self.entries[k].get().strip() if k in self.entries else ""
            
            from utils import calculate_depreciation
            
            new_depr = calculate_depreciation(
                valor=v("VALOR"),
                existencia_inicial=v("EXISTENCIAINICIAL"),
                descripcion=v("DESCRIPCION"),
                ubicacion=v("UBICACION"),
                fecha_adquisicion=v("FECHAADQUISICION"),
                vida_util=v("VIDAUTIL")
            )
            
            dep_entry = self.entries.get("DEPRECIACUMULADA")
            if dep_entry and self.focus_get() != dep_entry:
                dep_entry.delete(0, 'end')
                dep_entry.insert(0, str(new_depr))
        except Exception:
            pass

    def delete_current_asset(self):
        desc = self.asset.get('DESCRIPCION') or self.asset.get('description') or 'este activo'
        if messagebox.askyesno("Confirmar Eliminación", f"¿Está seguro de que desea eliminar definitivamente '{desc}' del inventario?\n\nEsta acción no se puede deshacer y recalculará los balances automáticamente."):
            asset_id = self.asset.get('id')
            if asset_id and self.parent.asset_repo.delete_asset(asset_id):
                messagebox.showinfo("Activo Eliminado", "El activo se ha eliminado correctamente del sistema.")
                self.on_success()
                self.destroy()
            else:
                messagebox.showerror("Error", "No se pudo eliminar el activo.")

    def create_section(self, title, icon):
        card = ctk.CTkFrame(self.scroll, fg_color="#0f172a", corner_radius=20, border_width=1, border_color="#1e293b")
        card.pack(fill="x", pady=10)
        
        h = ctk.CTkFrame(card, fg_color="#111827", height=40, corner_radius=0)
        h.pack(fill="x", padx=1, pady=1)
        ctk.CTkLabel(h, text=f"{icon}  {title}", font=ctk.CTkFont(size=12, weight="bold"), text_color="#60a5fa").pack(side="left", padx=15)
        
        grid = ctk.CTkFrame(card, fg_color="transparent")
        grid.pack(fill="x", padx=20, pady=15)
        grid.columnconfigure((0, 1), weight=1)
        return grid

    def create_field(self, parent, label, key, default_val, row, col, colspan=1, values=None):
        f = ctk.CTkFrame(parent, fg_color="transparent")
        f.grid(row=row, column=col, columnspan=colspan, sticky="nsew", padx=10, pady=8)
        
        ctk.CTkLabel(f, text=label.upper(), font=ctk.CTkFont(size=10, weight="bold"), text_color="#475569").pack(anchor="w", padx=5, pady=(0, 4))
        
        val = "" if default_val is None else str(default_val)
        if values:
            w = ctk.CTkComboBox(f, values=values, height=40, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
            w.pack(fill="x")
            w.set(val if val in values else values[0])
        else:
            w = ctk.CTkEntry(f, height=40, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
            w.pack(fill="x")
            w.insert(0, val)
        self.entries[key] = w

    def save(self):
        try:
            def v(k): return self.entries[k].get().strip() if k in self.entries else ""
            
            def validate_date(date_str):
                if not date_str: return True
                for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
                    try:
                        datetime.strptime(date_str, fmt)
                        return True
                    except ValueError:
                        pass
                return False

            def is_non_negative_float(val_str):
                if not val_str: return True
                try:
                    f = float(str(val_str).replace('$','').replace(',','').replace(' ',''))
                    return f >= 0.0
                except ValueError:
                    return False

            def is_non_negative_int(val_str):
                if not val_str: return True
                try:
                    i = int(str(val_str).replace(',','').replace(' ',''))
                    return i >= 0
                except ValueError:
                    return False

            def validate_ident(ident_str):
                if not ident_str: return True
                return bool(re.match(r'^[0-9a-zA-Z.-]+$', ident_str))

            desc = v("DESCRIPCION")
            if not desc:
                messagebox.showwarning("Campo Requerido", "La descripción es obligatoria.")
                return
                
            # Validar Fecha
            if v("FECHAADQUISICION") and not validate_date(v("FECHAADQUISICION")):
                messagebox.showwarning("Fecha Inválida", "La fecha de ingreso debe tener un formato válido (DD/MM/AAAA o AAAA-MM-DD).")
                return

            # Validar Valores Contables
            if v("VALOR") and not is_non_negative_float(v("VALOR")):
                messagebox.showwarning("Valor Inválido", "El valor debe ser un número no negativo.")
                return
            if v("DEPRECIACUMULADA") and not is_non_negative_float(v("DEPRECIACUMULADA")):
                messagebox.showwarning("Valor Inválido", "La depreciación acumulada debe ser un número no negativo.")
                return

            # Validar Existencia Inicial y Vida Útil
            if v("EXISTENCIAINICIAL") and not is_non_negative_int(v("EXISTENCIAINICIAL")):
                messagebox.showwarning("Cantidad Inválida", "La existencia inicial debe ser un número entero no negativo.")
                return
            if v("VIDAUTIL"):
                try:
                    life_val = int(v("VIDAUTIL"))
                    if life_val < 0: raise ValueError()
                except ValueError:
                    messagebox.showwarning("Vida Útil Inválida", "La vida útil debe ser un número entero no negativo.")
                    return

            # Validar Identificación/NIT
            if v("IDENTIFICACION") and not validate_ident(v("IDENTIFICACION")):
                messagebox.showwarning("Identificación Inválida", "La identificación / NIT debe contener únicamente números, letras, puntos o guiones.")
                return

            def clean_num(val):
                if not val: return 0.0
                try: return float(str(val).replace('$','').replace(',','').replace(' ',''))
                except: return 0.0
                
            val_cost = clean_num(v("VALOR"))
            val_depr = clean_num(v("DEPRECIACUMULADA"))
            qty = max(0, int(clean_num(v("EXISTENCIAINICIAL"))))
            try: life = int(clean_num(v("VIDAUTIL")))
            except: life = 10

            updated_data = {
                "DESCRIPCION": desc,
                "VALOR": val_cost,
                "FECHAADQUISICION": v("FECHAADQUISICION"),
                "MARCA": v("MARCA") or "N/A",
                "SERIAL": v("SERIAL") or "N/A",
                "UBICACION": v("UBICACION") or "SIN ASIGNAR",
                "CODCONTABLE": v("CODCONTABLE") or "",
                "VIDAUTIL": life,
                "CODDEPRECIACION": v("CODDEPRECIACION") or "",
                "DEPRECIACUMULADA": val_depr,
                "CODGASTO": v("CODGASTO") or "",
                "FUNCIONARIO": v("FUNCIONARIO") or "",
                "IDENTIFICACION": v("IDENTIFICACION") or "",
                "CODGRUPO": v("CODGRUPO") or "",
                "CODSUBGRUPO": v("CODSUBGRUPO") or "",
                "TIPO": v("TIPO") or "D",
                "EXISTENCIAINICIAL": qty,
                "SEDE": v("SEDE") or ""
            }
            
            success = self.parent.asset_repo.update_asset(self.asset["id"], updated_data)
            if success:
                messagebox.showinfo("Éxito", "Activo actualizado correctamente.")
                self.on_success()
                self.destroy()
            else:
                messagebox.showerror("Error", "No se pudo actualizar el activo en la base de datos.")
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al guardar los cambios: {e}")


class TransferAssetDialog(ctk.CTkToplevel):
    def __init__(self, parent, asset_dict, on_success_callback):
        super().__init__(parent)
        self.parent = parent
        self.asset = asset_dict
        self.on_success = on_success_callback
        
        self.title(f"TRASLADAR ACTIVO: {self.asset.get('code')}")
        self.geometry("500x350")
        self.configure(fg_color="#0a0f1d")
        self.resizable(False, False)
        
        self.grab_set()
        self.focus()
        self.transient(parent)

        # Center dialog
        self.update_idletasks()
        try:
            width, height = 500, 350
            x = parent.winfo_rootx() + (parent.winfo_width() // 2) - (width // 2)
            y = parent.winfo_rooty() + (parent.winfo_height() // 2) - (height // 2)
            self.geometry(f"{width}x{height}+{x}+{y}")
        except Exception:
            pass
        
        container = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=20, border_width=1, border_color="#1e293b")
        container.pack(fill="both", expand=True, padx=20, pady=20)
        
        # Header
        ctk.CTkLabel(container, text="TRASLADAR ACTIVO", font=ctk.CTkFont(size=11, weight="bold"), text_color="#3b82f6").pack(anchor="w", padx=25, pady=(20, 2))
        ctk.CTkLabel(container, text=f"{self.asset.get('description')[:50]}...", font=ctk.CTkFont(size=18, weight="bold"), text_color="#f8fafc").pack(anchor="w", padx=25, pady=(0, 15))
        
        # Current Location
        info_f = ctk.CTkFrame(container, fg_color="#111827", corner_radius=10)
        info_f.pack(fill="x", padx=25, pady=5)
        ctk.CTkLabel(info_f, text=f"Ubicación Actual:  {self.asset.get('dependency', 'NINGUNA')}", font=ctk.CTkFont(size=12, weight="bold"), text_color="#cbd5e1").pack(padx=15, pady=10, anchor="w")
        
        # Target Selector
        select_f = ctk.CTkFrame(container, fg_color="transparent")
        select_f.pack(fill="x", padx=25, pady=15)
        ctk.CTkLabel(select_f, text="SELECCIONAR NUEVO SALÓN / ÁREA", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        
        asset_sede = self.asset.get('SEDE') or self.asset.get('sede')
        db_deps = self.parent.dep_repo.get_dependencies(sede=asset_sede)
        available_deps = [d for d in db_deps if d.upper() != "GENERAL"] if db_deps else []
        if not available_deps:
            available_deps = ["SIN ASIGNAR"]
        
        self.dep_combo = ctk.CTkComboBox(select_f, values=available_deps, height=44, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.dep_combo.pack(fill="x")
        self.dep_combo.set(self.asset.get("dependency", available_deps[0]))
        
        # Actions
        actions = ctk.CTkFrame(container, fg_color="transparent")
        actions.pack(fill="x", padx=25, pady=(15, 20))
        
        ctk.CTkButton(actions, text="🔄 REALIZAR TRASLADO", font=ctk.CTkFont(size=13, weight="bold"), height=44, fg_color="#f59e0b", hover_color="#d97706", text_color="white", corner_radius=10, command=self.save).pack(side="right", fill="x", expand=True, padx=(8, 0))
        ctk.CTkButton(actions, text="Cancelar", font=ctk.CTkFont(size=13), height=44, fg_color="#1e293b", hover_color="#334155", corner_radius=10, command=self.destroy).pack(side="left", padx=(0, 8))

    def save(self):
        new_dep = self.dep_combo.get().strip()
        if not new_dep:
            messagebox.showwarning("Requerido", "Debe seleccionar un salón de destino.")
            return
            
        if new_dep == self.asset.get("dependency"):
            messagebox.showinfo("Información", "El activo ya se encuentra en esa dependencia.")
            self.destroy()
            return
            
        success = self.parent.asset_repo.transfer_asset(self.asset["id"], new_dep)
        if success:
            messagebox.showinfo("Éxito", f"Activo trasladado a {new_dep} correctamente.")
            self.on_success()
            self.destroy()
        else:
            messagebox.showerror("Error", "No se pudo realizar el traslado en la base de datos.")


class DisposeAssetDialog(ctk.CTkToplevel):
    def __init__(self, parent, asset_dict, on_success_callback):
        super().__init__(parent)
        self.parent = parent
        self.asset = asset_dict
        self.on_success = on_success_callback
        
        asset_code = self.asset.get('code') or self.asset.get('CODIGO') or ''
        asset_desc = self.asset.get('description') or self.asset.get('DESCRIPCION') or 'Activo'
        
        self.title(f"DAR DE BAJA ACTIVO: {asset_code}")
        self.configure(fg_color="#0a0f1d")
        self.resizable(False, False)
        
        self.grab_set()
        self.focus()
        self.transient(parent)
        
        # Center dialog
        self.update_idletasks()
        width, height = 550, 560
        x = parent.winfo_rootx() + (parent.winfo_width() // 2) - (width // 2)
        y = parent.winfo_rooty() + (parent.winfo_height() // 2) - (height // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")
        
        container = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=20, border_width=1, border_color="#1e293b")
        container.pack(fill="both", expand=True, padx=20, pady=20)
        
        # 1. Pinned bottom action buttons (ALWAYS visible)
        actions = ctk.CTkFrame(container, fg_color="transparent")
        actions.pack(side="bottom", fill="x", padx=20, pady=(10, 15))
        
        ctk.CTkButton(actions, text="📉 DECRETAR BAJA", font=ctk.CTkFont(size=13, weight="bold"), height=44, fg_color="#ef4444", hover_color="#dc2626", text_color="white", corner_radius=10, command=self.save).pack(side="right", fill="x", expand=True, padx=(8, 0))
        ctk.CTkButton(actions, text="Cancelar", font=ctk.CTkFont(size=13), height=44, fg_color="#1e293b", hover_color="#334155", corner_radius=10, command=self.destroy).pack(side="left", padx=(0, 8))

        # 2. Scrollable content area so nothing is clipped
        content = ctk.CTkScrollableFrame(container, fg_color="transparent")
        content.pack(fill="both", expand=True, padx=5, pady=(5, 5))

        # Header
        ctk.CTkLabel(content, text="DESINCORPORACIÓN DE ACTIVO", font=ctk.CTkFont(size=11, weight="bold"), text_color="#ef4444").pack(anchor="w", padx=15, pady=(10, 2))
        ctk.CTkLabel(content, text=f"{str(asset_desc)[:50]}...", font=ctk.CTkFont(size=18, weight="bold"), text_color="#f8fafc").pack(anchor="w", padx=15, pady=(0, 15))
        
        self.max_qty = int(self.asset.get("quantity") or self.asset.get("EXISTENCIAINICIAL") or 1)
        
        # Row 1
        row1 = ctk.CTkFrame(content, fg_color="transparent")
        row1.pack(fill="x", padx=15, pady=5)
        row1.columnconfigure(0, weight=1); row1.columnconfigure(1, weight=1)
        
        c1 = ctk.CTkFrame(row1, fg_color="transparent")
        c1.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        ctk.CTkLabel(c1, text=f"CANTIDAD (MÁX: {self.max_qty})", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        self.qty_entry = ctk.CTkEntry(c1, height=40, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.qty_entry.insert(0, str(self.max_qty))
        self.qty_entry.pack(fill="x")
        
        c2 = ctk.CTkFrame(row1, fg_color="transparent")
        c2.grid(row=0, column=1, sticky="ew", padx=(10, 0))
        ctk.CTkLabel(c2, text="FECHA DE BAJA (DD/MM/AAAA)", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        self.date_entry = ctk.CTkEntry(c2, height=40, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.date_entry.insert(0, datetime.now().strftime("%d/%m/%Y"))
        self.date_entry.pack(fill="x")
        
        # Summary
        self.summary_f = ctk.CTkFrame(content, fg_color="#020617", corner_radius=10, border_width=1, border_color="#1e293b")
        self.summary_f.pack(fill="x", padx=15, pady=(15, 0))
        
        self.lbl_initial = ctk.CTkLabel(self.summary_f, text="Valor Inicial: $0.00", font=ctk.CTkFont(size=11), text_color="#94a3b8")
        self.lbl_initial.pack(anchor="w", padx=15, pady=(10, 2))
        self.lbl_depr = ctk.CTkLabel(self.summary_f, text="Depreciación: $0.00", font=ctk.CTkFont(size=11), text_color="#ef4444")
        self.lbl_depr.pack(anchor="w", padx=15, pady=2)
        self.lbl_real = ctk.CTkLabel(self.summary_f, text="Valor Real (Desvalorización): $0.00", font=ctk.CTkFont(size=12, weight="bold"), text_color="#10b981")
        self.lbl_real.pack(anchor="w", padx=15, pady=(2, 10))
        
        self.qty_entry.bind("<KeyRelease>", self.update_summary)
        self.update_summary(None)
        
        # Reason
        reason_f = ctk.CTkFrame(content, fg_color="transparent")
        reason_f.pack(fill="x", padx=15, pady=15)
        ctk.CTkLabel(reason_f, text="MOTIVO / RESOLUCIÓN LEGAL DE LA BAJA", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        self.reason_entry = ctk.CTkEntry(reason_f, height=44, placeholder_text="Ej: Daño irreparable / Hurto / Obsolescencia...", corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.reason_entry.pack(fill="x")
        
        # Caution
        warning_f = ctk.CTkFrame(content, fg_color="#1e1b1b", corner_radius=10, border_width=1, border_color="#7f1d1d")
        warning_f.pack(fill="x", padx=15, pady=5)
        ctk.CTkLabel(warning_f, text="⚠️ Esta acción es irreversible. Se descontará del balance\ncontable y el historial se guardará en bajas.", font=ctk.CTkFont(size=11, weight="bold"), text_color="#f87171", justify="left").pack(padx=15, pady=10, anchor="w")

    def save(self):
        try:
            qty_raw = self.qty_entry.get().strip()
            reason = self.reason_entry.get().strip()
            date = self.date_entry.get().strip()
            
            if not qty_raw or not reason or not date:
                messagebox.showwarning("Requerido", "Todos los campos son obligatorios.")
                return
                
            # Validar fecha (Task 15)
            try:
                datetime.strptime(date, "%d/%m/%Y")
            except ValueError:
                messagebox.showwarning("Fecha Inválida", "La fecha de baja debe tener el formato DD/MM/AAAA.")
                return
                
            try:
                qty = int(qty_raw)
            except:
                messagebox.showwarning("Cantidad Inválida", "La cantidad debe ser un número entero.")
                return
                
            if qty <= 0:
                messagebox.showwarning("Cantidad Inválida", "La cantidad debe ser mayor a 0.")
                return
                
            if qty > self.max_qty:
                messagebox.showwarning("Cantidad Excedida", f"No puede dar de baja más unidades de las disponibles ({self.max_qty}).")
                return
                
            confirm = messagebox.askyesno("Confirmar Baja", f"¿Está seguro de que desea dar de baja {qty} unidades de este activo?")
            if not confirm:
                return
                
            success = self.parent.asset_repo.dispose_asset(self.asset["id"], reason, date, qty_to_dispose=qty)
            if success:
                messagebox.showinfo("Éxito", "Decreto de baja procesado correctamente.")
                self.on_success()
                self.destroy()
            else:
                messagebox.showerror("Error", "No se pudo procesar la baja en la base de datos.")
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al procesar la baja: {e}")

    def update_summary(self, event):
        try:
            qty = int(self.qty_entry.get().strip())
        except:
            qty = 0
            
        if qty < 0: qty = 0
        if qty > self.max_qty: qty = self.max_qty
        
        ratio = qty / self.max_qty if self.max_qty > 0 else 0
        
        val_unit = float(self.asset.get('initial_unit_value') or self.asset.get('VALOR') or 0.0)
        tot_val = float(self.asset.get('initial_total_value') or (val_unit * self.max_qty) or 0.0)
        initial_total = tot_val * ratio
        depr_total = float(self.asset.get('depreciation_rate') or self.asset.get('DEPRECIACUMULADA') or 0.0)
        depr = depr_total * ratio
        real = max(0.0, initial_total - depr)
        
        self.lbl_initial.configure(text=f"Valor Inicial ({qty} und): ${initial_total:,.2f}")
        self.lbl_depr.configure(text=f"Depreciación Acumulada: ${depr:,.2f}")
        self.lbl_real.configure(text=f"Valor Real (Desvalorización): ${real:,.2f}")


class BulkTransferDialog(ctk.CTkToplevel):
    def __init__(self, parent, asset_ids, on_success_callback):
        super().__init__(parent)
        self.parent = parent
        self.asset_ids = asset_ids
        self.on_success = on_success_callback
        
        self.title("TRASLADO MASIVO EN LOTE")
        self.geometry("500x350")
        self.configure(fg_color="#0a0f1d")
        self.resizable(False, False)
        
        self.grab_set()
        self.focus()
        self.transient(parent)
        
        # Center dialog
        self.update_idletasks()
        width, height = 500, 350
        x = parent.winfo_rootx() + (parent.winfo_width() // 2) - (width // 2)
        y = parent.winfo_rooty() + (parent.winfo_height() // 2) - (height // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")
        
        container = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=20, border_width=1, border_color="#1e293b")
        container.pack(fill="both", expand=True, padx=20, pady=20)
        
        ctk.CTkLabel(container, text="TRASLADO MASIVO", font=ctk.CTkFont(size=11, weight="bold"), text_color="#f59e0b").pack(anchor="w", padx=25, pady=(20, 2))
        ctk.CTkLabel(container, text=f"Se trasladarán {len(asset_ids)} activos seleccionados.", font=ctk.CTkFont(size=14), text_color="#cbd5e1").pack(anchor="w", padx=25, pady=(0, 15))
        
        # Target Selector
        select_f = ctk.CTkFrame(container, fg_color="transparent")
        select_f.pack(fill="x", padx=25, pady=15)
        ctk.CTkLabel(select_f, text="SELECCIONAR NUEVO SALÓN / ÁREA", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        
        current_sede = getattr(self.parent, 'current_sede', None)
        db_deps = self.parent.dep_repo.get_dependencies(sede=current_sede)
        available_deps = [d for d in db_deps if d.upper() != "GENERAL"] if db_deps else []
        if not available_deps:
            available_deps = ["SIN ASIGNAR"]
        
        self.dep_combo = ctk.CTkComboBox(select_f, values=available_deps, height=44, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.dep_combo.pack(fill="x")
        self.dep_combo.set(available_deps[0])
        
        actions = ctk.CTkFrame(container, fg_color="transparent")
        actions.pack(fill="x", padx=25, pady=(15, 20))
        
        ctk.CTkButton(actions, text="🔄 REALIZAR TRASLADO", font=ctk.CTkFont(size=13, weight="bold"), height=44, fg_color="#f59e0b", hover_color="#d97706", text_color="white", corner_radius=10, command=self.save).pack(side="right", fill="x", expand=True, padx=(8, 0))
        ctk.CTkButton(actions, text="Cancelar", font=ctk.CTkFont(size=13), height=44, fg_color="#1e293b", hover_color="#334155", corner_radius=10, command=self.destroy).pack(side="left", padx=(0, 8))

    def save(self):
        new_dep = self.dep_combo.get().strip()
        if not new_dep:
            messagebox.showwarning("Requerido", "Debe seleccionar un salón de destino.")
            return
            
        success_count = 0
        for aid in self.asset_ids:
            if self.parent.asset_repo.transfer_asset(aid, new_dep):
                success_count += 1
                
        if success_count > 0:
            messagebox.showinfo("Éxito", f"Se trasladaron {success_count} activos a {new_dep} correctamente.")
            self.on_success()
            self.destroy()
        else:
            messagebox.showerror("Error", "No se pudo realizar el traslado masivo.")


class BulkDisposeDialog(ctk.CTkToplevel):
    def __init__(self, parent, asset_ids, on_success_callback):
        super().__init__(parent)
        self.parent = parent
        self.asset_ids = asset_ids
        self.on_success = on_success_callback
        
        self.title("BAJA MASIVA EN LOTE")
        self.geometry("550x560")
        self.configure(fg_color="#0a0f1d")
        self.resizable(False, False)
        
        self.grab_set()
        self.focus()
        self.transient(parent)
        
        # Center dialog
        self.update_idletasks()
        width, height = 550, 560
        x = parent.winfo_rootx() + (parent.winfo_width() // 2) - (width // 2)
        y = parent.winfo_rooty() + (parent.winfo_height() // 2) - (height // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")
        
        container = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=20, border_width=1, border_color="#1e293b")
        container.pack(fill="both", expand=True, padx=20, pady=20)
        
        # 1. Pinned bottom action buttons (ALWAYS visible)
        actions = ctk.CTkFrame(container, fg_color="transparent")
        actions.pack(side="bottom", fill="x", padx=20, pady=(10, 15))
        
        ctk.CTkButton(actions, text="📉 DECRETAR BAJA EN LOTE", font=ctk.CTkFont(size=13, weight="bold"), height=44, fg_color="#ef4444", hover_color="#dc2626", text_color="white", corner_radius=10, command=self.save).pack(side="right", fill="x", expand=True, padx=(8, 0))
        ctk.CTkButton(actions, text="Cancelar", font=ctk.CTkFont(size=13), height=44, fg_color="#1e293b", hover_color="#334155", corner_radius=10, command=self.destroy).pack(side="left", padx=(0, 8))

        # 2. Scrollable content area so nothing is clipped
        content = ctk.CTkScrollableFrame(container, fg_color="transparent")
        content.pack(fill="both", expand=True, padx=5, pady=(5, 5))

        ctk.CTkLabel(content, text="DESINCORPORACIÓN MASIVA EN LOTE", font=ctk.CTkFont(size=11, weight="bold"), text_color="#ef4444").pack(anchor="w", padx=15, pady=(10, 2))
        ctk.CTkLabel(content, text=f"Dará de baja {len(asset_ids)} activos seleccionados.", font=ctk.CTkFont(size=14), text_color="#cbd5e1").pack(anchor="w", padx=15, pady=(0, 15))
        
        # Row 1
        row1 = ctk.CTkFrame(content, fg_color="transparent")
        row1.pack(fill="x", padx=15, pady=5)
        row1.columnconfigure(0, weight=1)
        
        c2 = ctk.CTkFrame(row1, fg_color="transparent")
        c2.grid(row=0, column=0, sticky="ew")
        ctk.CTkLabel(c2, text="FECHA DE BAJA (DD/MM/AAAA)", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        self.date_entry = ctk.CTkEntry(c2, height=40, corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.date_entry.insert(0, datetime.now().strftime("%d/%m/%Y"))
        self.date_entry.pack(fill="x")
        
        # Calculate summary for bulk
        total_initial = 0.0
        total_depr = 0.0
        for aid in self.asset_ids:
            try:
                a_data = self.parent.asset_repo.get_asset_by_id(aid)
                if a_data:
                    val_u = float(a_data.get('initial_unit_value') or a_data.get('VALOR') or 0.0)
                    q = int(a_data.get('quantity') or a_data.get('EXISTENCIAINICIAL') or 1)
                    tot = float(a_data.get('initial_total_value') or (val_u * q) or 0.0)
                    dep = float(a_data.get('depreciation_rate') or a_data.get('DEPRECIACUMULADA') or 0.0)
                    total_initial += tot
                    total_depr += dep
            except Exception as e:
                print(f"Error calculating bulk summary for asset {aid}: {e}")
        total_real = max(0.0, total_initial - total_depr)
        
        # Summary
        self.summary_f = ctk.CTkFrame(content, fg_color="#020617", corner_radius=10, border_width=1, border_color="#1e293b")
        self.summary_f.pack(fill="x", padx=15, pady=(15, 0))
        
        ctk.CTkLabel(self.summary_f, text=f"Valor Inicial Total: ${total_initial:,.2f}", font=ctk.CTkFont(size=11), text_color="#94a3b8").pack(anchor="w", padx=15, pady=(10, 2))
        ctk.CTkLabel(self.summary_f, text=f"Depreciación Acumulada: ${total_depr:,.2f}", font=ctk.CTkFont(size=11), text_color="#ef4444").pack(anchor="w", padx=15, pady=2)
        ctk.CTkLabel(self.summary_f, text=f"Valor Real (Desvalorización): ${total_real:,.2f}", font=ctk.CTkFont(size=12, weight="bold"), text_color="#10b981").pack(anchor="w", padx=15, pady=(2, 10))
        
        # Reason
        reason_f = ctk.CTkFrame(content, fg_color="transparent")
        reason_f.pack(fill="x", padx=15, pady=15)
        ctk.CTkLabel(reason_f, text="MOTIVO / RESOLUCIÓN LEGAL DE LA BAJA", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", pady=(0, 5))
        self.reason_entry = ctk.CTkEntry(reason_f, height=44, placeholder_text="Ej: Daño irreparable / Hurto / Obsolescencia...", corner_radius=10, border_width=1, border_color="#1e293b", fg_color="#020617")
        self.reason_entry.pack(fill="x")
        
        # Caution
        warning_f = ctk.CTkFrame(content, fg_color="#1e1b1b", corner_radius=10, border_width=1, border_color="#7f1d1d")
        warning_f.pack(fill="x", padx=15, pady=5)
        ctk.CTkLabel(warning_f, text="⚠️ Esta acción dará de baja la TOTALIDAD de existencias de\ntodos los activos seleccionados. Acción irreversible.", font=ctk.CTkFont(size=11, weight="bold"), text_color="#f87171", justify="left").pack(padx=15, pady=10, anchor="w")

    def save(self):
        reason = self.reason_entry.get().strip()
        date = self.date_entry.get().strip()
        
        if not reason or not date:
            messagebox.showwarning("Requerido", "Todos los campos son obligatorios.")
            return
            
        # Validar fecha (Task 15)
        try:
            datetime.strptime(date, "%d/%m/%Y")
        except ValueError:
            messagebox.showwarning("Fecha Inválida", "La fecha de baja debe tener el formato DD/MM/AAAA.")
            return
            
        confirm = messagebox.askyesno("Confirmar Baja Masiva", f"¿Está seguro de que desea dar de baja los {len(self.asset_ids)} activos seleccionados?")
        if not confirm:
            return
            
        success_count = 0
        for aid in self.asset_ids:
            if self.parent.asset_repo.dispose_asset(aid, reason, date):
                success_count += 1
                
        if success_count > 0:
            messagebox.showinfo("Éxito", f"Decreto de baja procesado para {success_count} activos.")
            self.on_success()
            self.destroy()
        else:
            messagebox.showerror("Error", "No se pudo realizar la baja masiva.")
