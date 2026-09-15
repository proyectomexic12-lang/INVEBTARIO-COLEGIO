import os
import sys
import logger_config  # <-- Logger diagnostic system initialized globally here

# Configure isolated high-speed cache for matplotlib to eradicate OS lock latency
import tempfile
_mpl_dir = os.path.join(tempfile.gettempdir(), "inventario_mplcache")
os.makedirs(_mpl_dir, exist_ok=True)
os.environ["MPLCONFIGDIR"] = _mpl_dir

import customtkinter as ctk
import tkinter
import openpyxl
import openpyxl.reader.excel
import fpdf

# Safely suppress Tcl/Tkinter teardown errors during Python GC and exit
try:
    _orig_var_del = tkinter.Variable.__del__
    def _safe_var_del(self):
        try:
            _orig_var_del(self)
        except Exception:
            pass
    tkinter.Variable.__del__ = _safe_var_del

    _orig_img_del = tkinter.Image.__del__
    def _safe_img_del(self):
        try:
            _orig_img_del(self)
        except Exception:
            pass
    tkinter.Image.__del__ = _safe_img_del
except Exception:
    pass

import threading
import queue
from database import DatabaseManager, AuditRepository, SettingRepository, DependencyRepository, AssetRepository, LoanRepository, OfficialRepository
from theme import theme

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


class LoadingDialog(ctk.CTkToplevel):
    def __init__(self, parent, message="Procesando..."):
        super().__init__(parent)
        self.title("Por favor espere")
        self.geometry("320x130")
        self.configure(fg_color="#0a0f1d")
        self.resizable(False, False)
        
        try:
            self.grab_set()
        except Exception:
            pass
        self.focus()
        self.transient(parent)
        
        self.update_idletasks()
        width, height = 320, 130
        x = parent.winfo_rootx() + (parent.winfo_width() // 2) - (width // 2)
        y = parent.winfo_rooty() + (parent.winfo_height() // 2) - (height // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")
        
        self.protocol("WM_DELETE_WINDOW", self.close_safe)
        
        container = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=15, border_width=1, border_color="#1e293b")
        container.pack(fill="both", expand=True, padx=10, pady=10)
        
        self.lbl = ctk.CTkLabel(container, text=message, font=ctk.CTkFont(size=13, weight="bold"), text_color="#3b82f6", wraplength=280)
        self.lbl.pack(pady=(15, 10))
        
        self.pbar = ctk.CTkProgressBar(container, width=220, height=8, corner_radius=4, progress_color="#3b82f6", fg_color="#1e293b")
        self.pbar.pack(pady=(0, 15))
        self.pbar.configure(mode="indefinite")
        self.pbar.start()

    def close_safe(self):
        try:
            self.grab_release()
        except Exception:
            pass
        self.destroy()

    def destroy(self):
        try:
            self.grab_release()
        except Exception:
            pass
        super().destroy()

class SafeModeWindow(ctk.CTk):
    def __init__(self, error_message):
        super().__init__()
        self.title("MODO DE RECUPERACIÓN DE EMERGENCIA")
        self.geometry("600x400")
        self.configure(fg_color="#0a0f1d")
        
        lbl = ctk.CTkLabel(self, text="⚠️ ERROR CRÍTICO AL INICIAR", font=ctk.CTkFont(size=24, weight="bold"), text_color="#ef4444")
        lbl.pack(pady=(40, 10))
        
        msg = ctk.CTkLabel(self, text="La base de datos principal está corrupta o es inaccesible.\nEl sistema no puede iniciar normalmente para proteger los datos.", 
                           font=ctk.CTkFont(size=14))
        msg.pack(pady=10)
        
        err = ctk.CTkTextbox(self, width=500, height=100, text_color="#f87171", fg_color="#1e293b")
        err.pack(pady=10)
        err.insert("1.0", str(error_message))
        err.configure(state="disabled")
        
        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.pack(pady=20)
        
        btn_export = ctk.CTkButton(btn_frame, text="Extraer Backup de Emergencia", fg_color="#3b82f6", hover_color="#2563eb",
                                   command=self.extract_backup)
        btn_export.pack(side="left", padx=10)
        
        btn_exit = ctk.CTkButton(btn_frame, text="Salir", fg_color="#ef4444", hover_color="#dc2626", command=self.destroy)
        btn_exit.pack(side="left", padx=10)

    def extract_backup(self):
        import shutil
        from tkinter import filedialog, messagebox
        db_path = os.path.join(os.path.dirname(__file__), 'inventario_guaimaral.db')
        target = filedialog.asksaveasfilename(defaultextension=".db", initialfile="backup_corrupto.db", title="Guardar DB para análisis")
        if target:
            try:
                shutil.copy2(db_path, target)
                messagebox.showinfo("Éxito", f"Base de datos exportada a:\n{target}\n\nPor favor, contacte a soporte técnico.")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo copiar: {e}")

class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        # Initialize Database Manager and Repositories (Composition Root)
        try:
            self.db_manager = DatabaseManager()
        except Exception as e:
            import logger_config
            logger_config.global_logger.critical("Error fatal al inicializar base de datos: %s", e)
            self.destroy()
            safe_app = SafeModeWindow(e)
            safe_app.mainloop()
            sys.exit(1)

        self.audit_repo = AuditRepository(self.db_manager)
        self.setting_repo = SettingRepository(self.db_manager)
        self.dep_repo = DependencyRepository(self.db_manager)
        self.asset_repo = AssetRepository(self.db_manager, self.setting_repo, self.audit_repo)
        self.loan_repo = LoanRepository(self.db_manager)
        self.official_repo = OfficialRepository(self.db_manager)
        
        # Load theme setting from DB first
        try:
            theme_name = self.setting_repo.get_setting("tema", "Obsidian Executive")
            if theme_name:
                theme.set_theme_by_name(theme_name)
        except Exception as e:
            print(f"[INVENTARIO] Error cargando tema de DB: {e}")

        # Ejecutar copia de seguridad automática (Task 18)
        threading.Thread(target=self.backup_database, daemon=True).start()

        self.title("CONTROL INVENTARIO I.E. GUAIMARAL")
        self.geometry("1450x950")
        self.configure(fg_color=theme.BG_APP)
        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=1)

        self.protocol("WM_DELETE_WINDOW", self.on_closing)

        # Track current page
        self.current_page_fk = "ModernDashboard"
        self.current_page_text = "Panel Maestro"

        # =====================================================
        # 1. INFRAESTRUCTURA (PRIMERO QUE TODO)
        # =====================================================
        self.queue = queue.Queue()
        self.process_queue()

        # =====================================================
        # 2. SIDEBAR
        # =====================================================
        self.sidebar = ctk.CTkFrame(self, width=320, fg_color=theme.BG_SIDEBAR, corner_radius=0)
        self.sidebar.grid(row=0, column=0, sticky="nsew")

        brand_f = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        brand_f.pack(pady=(40, 30), padx=30)

        self.logo_icon = ctk.CTkFrame(brand_f, width=50, height=50, corner_radius=15, fg_color=theme.PRIMARY)
        self.logo_icon.pack(side="left")
        self.logo_icon.pack_propagate(False)
        ctk.CTkLabel(self.logo_icon, text="IEG", font=ctk.CTkFont(size=18, weight="bold"), text_color="white").pack(expand=True)

        logo_txt_f = ctk.CTkFrame(brand_f, fg_color="transparent")
        logo_txt_f.pack(side="left", padx=15)
        self.logo_txt = ctk.CTkLabel(logo_txt_f, text="INVENTARIO", font=ctk.CTkFont(size=22, weight="bold"), text_color=theme.TEXT_MAIN)
        self.logo_txt.pack(anchor="w")
        self.logo_sub = ctk.CTkLabel(logo_txt_f, text="I.E. GUAIMARAL", font=ctk.CTkFont(size=11, weight="bold"), text_color=theme.PRIMARY)
        self.logo_sub.pack(anchor="w")

        # Barra de búsqueda global
        self.search_f = ctk.CTkFrame(self.sidebar, fg_color=theme.BG_INPUT, corner_radius=15, height=45, border_width=1, border_color=theme.BORDER)
        self.search_f.pack(fill="x", padx=20, pady=(0, 20))
        self.search_f.pack_propagate(False)
        
        self.search_filter = ctk.CTkComboBox(self.search_f, values=["Todo", "Código", "Descripción", "Funcionario"],
                                             width=85, height=28, corner_radius=10, border_width=0,
                                             fg_color=theme.BG_APP, button_color=theme.BORDER,
                                             command=self.perform_instant_search)
        self.search_filter.set("Todo")
        self.search_filter.pack(side="right", padx=5)

        ctk.CTkLabel(self.search_f, text="🔍", font=ctk.CTkFont(size=16)).pack(side="left", padx=(12, 4))
        self.search_entry = ctk.CTkEntry(self.search_f, placeholder_text="Búsqueda rápida...",
                                         fg_color="transparent", border_width=0,
                                         font=ctk.CTkFont(size=13))
        self.search_entry.pack(side="left", fill="both", expand=True)
        self.search_entry.bind("<KeyRelease>", self.perform_instant_search)

        # Sede Selector Frame (Task 17)
        self.sede_f = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        self.sede_f.pack(fill="x", padx=20, pady=(0, 20))
        ctk.CTkLabel(self.sede_f, text="🏫 SEDE ACTIVA GLOBAL", font=ctk.CTkFont(size=10, weight="bold"), text_color=theme.TEXT_MUTED).pack(anchor="w", padx=5, pady=(0, 4))
        
        self.available_sedes = ["Todas"] + self.asset_repo.get_all_sedes()
        self.sede_combo = ctk.CTkComboBox(self.sede_f, values=self.available_sedes, height=40, corner_radius=12,
                                          fg_color=theme.BG_INPUT, border_color=theme.BORDER, button_color=theme.BORDER,
                                          command=self.change_global_sede)
        self.sede_combo.pack(fill="x")
        self.sede_combo.set("Todas")
        self.global_sede = "Todas"

        self.nav_items = [
            ("Panel Maestro",            "📊", "ModernDashboard"),
            ("Unidad de Registro",       "📦", "ModernRegistry"),
            ("Gestión Global",           "🌍", "ModernInventory"),
            ("Préstamos y Asignaciones", "🤝", "ModernLoans"),
            ("Salones y Áreas",          "🏢", "ModernDependencies"),
            ("Gestión del Sistema",      "⚙️", "ModernSettings"),
        ]

        self.buttons = {}
        self.btn_active_light = {}

        for text, icon, frame_key in self.nav_items:
            container = ctk.CTkFrame(self.sidebar, fg_color="transparent")
            container.pack(fill="x", padx=15, pady=2)

            light = ctk.CTkFrame(container, width=5, height=35, fg_color="transparent", corner_radius=10)
            light.pack(side="left")
            self.btn_active_light[text] = light

            btn = ctk.CTkButton(
                container, text=f"  {icon}   {text}",
                font=ctk.CTkFont(size=14, weight="bold"),
                height=48, corner_radius=12,
                fg_color="transparent", text_color=theme.TEXT_MUTED,
                anchor="w", hover_color=theme.BORDER,
                command=lambda fk=frame_key, t=text: self.select_page(fk, t)
            )
            btn.pack(side="left", fill="x", expand=True, padx=(8, 0))
            self.buttons[text] = btn

        self.profile_f = ctk.CTkFrame(self.sidebar, fg_color=theme.BG_INPUT, corner_radius=20,
                               border_width=1, border_color=theme.BORDER)
        self.profile_f.pack(side="bottom", fill="x", padx=20, pady=30)
        self.profile_lbl = ctk.CTkLabel(self.profile_f, text="⚡ SISTEMA ACTIVO",
                     font=ctk.CTkFont(size=9, weight="bold"), text_color=theme.SUCCESS)
        self.profile_lbl.pack(pady=(12, 12))

        # =====================================================
        # 3. WORKSPACE + PÁGINAS
        # =====================================================
        self.workspace = ctk.CTkFrame(self, fg_color="transparent")
        self.workspace.grid(row=0, column=1, sticky="nsew", padx=20, pady=20)
        self.workspace.grid_rowconfigure(0, weight=1)
        self.workspace.grid_columnconfigure(0, weight=1)

        # Instant Launch: Lazy Loading Architecture (SOLID - High Performance)
        def _get_dashboard():
            from components.dashboard import ModernDashboard
            return ModernDashboard(self.workspace, self, self.asset_repo)

        def _get_registry():
            from components.registry import ModernRegistry
            return ModernRegistry(self.workspace, self, self.asset_repo, self.dep_repo, self.setting_repo)

        def _get_inventory():
            from components.inventory import ModernInventory
            return ModernInventory(self.workspace, self, self.asset_repo, self.dep_repo)

        def _get_dependencies():
            from components.dependencies import ModernDependencies
            return ModernDependencies(self.workspace, self, self.dep_repo, self.asset_repo, self.setting_repo)

        def _get_loans():
            from components.loans import ModernLoans
            return ModernLoans(self.workspace, self, self.loan_repo, self.asset_repo, self.official_repo)

        def _get_settings():
            from components.settings import ModernSettings
            return ModernSettings(self.workspace, self, self.setting_repo, self.audit_repo)

        self.page_factories = {
            "ModernDashboard":    _get_dashboard,
            "ModernRegistry":     _get_registry,
            "ModernInventory":    _get_inventory,
            "ModernDependencies": _get_dependencies,
            "ModernLoans":        _get_loans,
            "ModernSettings":     _get_settings,
        }
        self.pages = {}
        self.needs_refresh = {k: True for k in self.page_factories}

        # Render window and sidebar shell immediately on screen (Zero perceptual delay)
        self.update()

        self.select_page("ModernDashboard", "Panel Maestro")

    def on_closing(self):
        import os
        try:
            self.quit()
            self.destroy()
        except:
            pass
        os._exit(0)

    def backup_database(self):
        import shutil
        from datetime import datetime
        import os
        import time
        try:
            time.sleep(1.2)  # Yield CPU during initial GUI render
            db_path = self.db_manager.db_path
            if not os.path.exists(db_path):
                return
            backup_dir = os.path.join(os.path.dirname(db_path), "backups")
            if not os.path.exists(backup_dir):
                os.makedirs(backup_dir)
                
            # Solo hacer backup si el último fue hace más de 1 hora
            existing_backups = [os.path.join(backup_dir, f) for f in os.listdir(backup_dir) if f.startswith("inventory_") and f.endswith(".db")]
            if existing_backups:
                newest = max(existing_backups, key=os.path.getmtime)
                if time.time() - os.path.getmtime(newest) < 3600:
                    return # No hacer backup todavía, ha pasado menos de 1 hora
                    
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_name = f"inventory_{timestamp}.db"
            backup_path = os.path.join(backup_dir, backup_name)
            self.db_manager.backup_to_file(backup_path)
            print(f"[INVENTARIO] Backup atómico y consistente creado en: {backup_path}")
            
            # Keep only the last 5 backups
            backups = sorted(
                [os.path.join(backup_dir, f) for f in os.listdir(backup_dir) if f.startswith("inventory_") and f.endswith(".db")],
                key=os.path.getmtime
            )
            while len(backups) > 5:
                oldest = backups.pop(0)
                try:
                    os.remove(oldest)
                except Exception:
                    pass
        except Exception as e:
            print(f"[INVENTARIO] Error al crear copia de seguridad: {e}")

    # =========================================================
    # THREADING
    # =========================================================
    def run_in_thread(self, task_func, callback=None, *args, **kwargs):
        def worker():
            try:
                result = task_func(*args, **kwargs)
                self.queue.put(("CALLBACK", callback, result))
            except Exception as e:
                import traceback
                traceback.print_exc()
                self.queue.put(("ERROR", str(e), None))
        threading.Thread(target=worker, daemon=True).start()

    def run_in_thread_with_loader(self, loader_text, task_func, callback=None, *args, **kwargs):
        loader = LoadingDialog(self, loader_text)
        self.active_loader = loader
        
        def safe_close_loader():
            if hasattr(self, "active_loader") and self.active_loader:
                try:
                    self.active_loader.grab_release()
                except Exception:
                    pass
                try:
                    self.active_loader.destroy()
                except Exception:
                    pass
                self.active_loader = None
            try:
                loader.grab_release()
            except Exception:
                pass
            try:
                loader.destroy()
            except Exception:
                pass

        def wrapped_callback(result):
            safe_close_loader()
            if callback:
                callback(result)

        def wrapped_error(err_msg):
            safe_close_loader()
            from tkinter import messagebox
            messagebox.showerror("Aviso del Sistema", f"{err_msg}")

        def worker():
            try:
                res = task_func(*args, **kwargs)
                self.queue.put(("CALLBACK", wrapped_callback, res))
            except Exception as e:
                import traceback
                traceback.print_exc()
                self.queue.put(("CALLBACK", wrapped_error, str(e)))

        threading.Thread(target=worker, daemon=True).start()

    def process_queue(self):
        try:
            while True:
                item = self.queue.get_nowait()
                msg_type = item[0]
                if msg_type == "CALLBACK":
                    _, func, data = item
                    if func:
                        try:
                            func(data)
                        except Exception as e:
                            err_str = str(e).lower()
                            if "bad window path name" not in err_str:
                                print(f"[INVENTARIO] Error en callback: {e}")
                elif msg_type == "ERROR":
                    _, msg, _ = item
                    print(f"[INVENTARIO] Error en hilo: {msg}")
                    # Release any grabbed modal dialogs
                    if hasattr(self, "active_loader") and self.active_loader:
                        try:
                            self.active_loader.grab_release()
                        except Exception:
                            pass
                        try:
                            self.active_loader.destroy()
                        except Exception:
                            pass
                        self.active_loader = None
                    for child in self.winfo_children():
                        if "loadingdialog" in str(type(child)).lower():
                            try:
                                child.grab_release()
                            except Exception:
                                pass
                            try:
                                child.destroy()
                            except Exception:
                                pass
                    from tkinter import messagebox
                    messagebox.showerror("Aviso del Sistema", f"{msg}")
        except queue.Empty:
            pass
        self.after(20, self.process_queue)

    # =========================================================
    # NAVEGACIÓN
    # =========================================================
    def select_page(self, name, btn_text):
        # Update current page tracking
        self.current_page_fk = name
        self.current_page_text = btn_text

        # Resetear botones
        for b in self.buttons.values():
            b.configure(fg_color="transparent", text_color=theme.TEXT_MUTED)
        for l in self.btn_active_light.values():
            l.configure(fg_color="transparent")

        if btn_text in self.buttons:
            self.buttons[btn_text].configure(fg_color=theme.BORDER, text_color=theme.PRIMARY)
            self.btn_active_light[btn_text].configure(fg_color=theme.PRIMARY)

        # Lazy Instantiation: load page on demand
        if name not in self.pages or self.pages[name] is None:
            if name in self.page_factories:
                self.pages[name] = self.page_factories[name]()

        # Mostrar página activa ocultando las demás creadas
        for p_k, p in self.pages.items():
            if p is not None and p_k != name:
                p.grid_forget()

        if name in self.pages and self.pages[name] is not None:
            self.pages[name].grid(row=0, column=0, sticky="nsew")

            # Solo refrescar datos si la página tiene cambios pendientes (Zero Lag)
            if self.needs_refresh.get(name, True):
                p_obj = self.pages[name]
                for method in ("refresh_data", "refresh_stats", "refresh_grid", "refresh_list", "render_main_grid", "refresh_form"):
                    if hasattr(p_obj, method):
                        getattr(p_obj, method)()
                        break   # solo llamar el primero que exista
                self.needs_refresh[name] = False

    def change_global_sede(self, selected_sede):
        self.global_sede = selected_sede
        self.invalidate_all_pages()
        self.select_page(self.current_page_fk, self.current_page_text)

    def update_sede_combobox(self):
        if hasattr(self, "sede_combo"):
            self.available_sedes = ["Todas"] + self.asset_repo.get_all_sedes()
            self.sede_combo.configure(values=self.available_sedes)
            if self.global_sede not in self.available_sedes:
                self.global_sede = "Todas"
            self.sede_combo.set(self.global_sede)

    def invalidate_all_pages(self):
        self.needs_refresh = {k: True for k in self.page_factories}
        self.update_sede_combobox()
        # High performance: only refresh the CURRENT visible page immediately.
        # Other pages will refresh automatically when clicked without freezing the UI.
        current_obj = self.pages.get(self.current_page_fk)
        if current_obj is not None:
            for method in ("refresh_data", "refresh_stats", "refresh_grid", "refresh_list", "render_main_grid", "refresh_form"):
                if hasattr(current_obj, method):
                    try:
                        getattr(current_obj, method)()
                    except Exception:
                        pass
                    break
            self.needs_refresh[self.current_page_fk] = False

    def recreate_all_pages(self):
        # Re-apply window background
        self.configure(fg_color=theme.BG_APP)
        
        # Re-configure sidebar and brand
        self.sidebar.configure(fg_color=theme.BG_SIDEBAR)
        self.logo_icon.configure(fg_color=theme.PRIMARY)
        self.logo_txt.configure(text_color=theme.TEXT_MAIN)
        self.logo_sub.configure(text_color=theme.PRIMARY)
        self.search_f.configure(fg_color=theme.BG_INPUT, border_color=theme.BORDER)
        
        # Re-configure navigation buttons
        for text, icon, frame_key in self.nav_items:
            if text in self.buttons:
                self.buttons[text].configure(text_color=theme.TEXT_MUTED, hover_color=theme.BORDER)
                
        self.profile_f.configure(fg_color=theme.BG_INPUT, border_color=theme.BORDER)
        self.profile_lbl.configure(text_color=theme.SUCCESS)
        
        # Clean and destroy instantiated pages
        for p in self.pages.values():
            if p is not None:
                try:
                    p.destroy()
                except Exception:
                    pass
            
        self.pages = {}
        self.needs_refresh = {k: True for k in self.page_factories}
        
        # Redraw current page with new colors
        self.select_page(self.current_page_fk, self.current_page_text)

    def sidebar_button_event(self, name):
        nav_map = {t: f for t, _, f in self.nav_items}
        if name in nav_map:
            self.select_page(nav_map[name], name)

    # =========================================================
    # BÚSQUEDA GLOBAL (CON DEBOUNCE ULTRA-RÁPIDO)
    # =========================================================
    def perform_instant_search(self, event=None):
        if hasattr(self, "_search_timer") and self._search_timer:
            try:
                self.after_cancel(self._search_timer)
            except Exception:
                pass

        def _do_search():
            query = self.search_entry.get().strip()
            filt = self.search_filter.get()
            if len(query) > 1:
                self.run_in_thread(self.asset_repo.instant_search, self._show_search_results, query, filt)
            elif len(query) == 0:
                self.select_page("ModernDashboard", "Panel Maestro")

        self._search_timer = self.after(200, _do_search)

    def _show_search_results(self, results):
        if not results:
            return
        self.select_page("ModernInventory", "Gestión Global")
        inv = self.pages.get("ModernInventory")
        if inv and hasattr(inv, "inject_results"):
            inv.inject_results(results)

    # =========================================================
    # IMPORTACIÓN INTELIGENTE
    # =========================================================
    def run_intelligent_import(self, filepath, target_dep=None):
        """Motor de importación de alta fidelidad con mapeo automático de columnas y fechas."""
        import openpyxl
        import re
        import unicodedata
        import os
        from datetime import datetime, date, timedelta

        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".xls":
            raise Exception("El archivo seleccionado está en formato antiguo de Excel (.xls).\n\nopenpyxl requiere formato moderno (.xlsx).\nPor favor, abra el archivo en Excel y seleccione 'Guardar como' -> 'Libro de Excel (*.xlsx)'.")

        try:
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        except PermissionError:
            # Si el archivo está abierto en Microsoft Excel, copiar a temporal y leerlo
            import tempfile
            import shutil
            try:
                temp_dir = tempfile.gettempdir()
                temp_copy = os.path.join(temp_dir, f"temp_import_{os.path.basename(filepath)}")
                shutil.copyfile(filepath, temp_copy)
                wb = openpyxl.load_workbook(temp_copy, read_only=True, data_only=True)
            except Exception:
                raise Exception("El archivo Excel está abierto y bloqueado por Excel.\n\nPor favor guarde y CIERRE el archivo en Excel antes de importarlo.")
        except Exception as e:
            raise Exception(f"No se pudo abrir el archivo Excel:\n{e}")

        sheet = wb.active
        # Si hay varias hojas y la activa parece vacía, buscar una hoja con filas de datos
        if len(wb.sheetnames) > 1:
            has_rows = False
            for r in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
                if any(r):
                    has_rows = True
                    break
            if not has_rows:
                for sname in wb.sheetnames:
                    candidate = wb[sname]
                    for r in candidate.iter_rows(min_row=1, max_row=5, values_only=True):
                        if any(r):
                            sheet = candidate
                            break
                    if sheet == candidate:
                        break

        def strip_accents(text):
            return ''.join(c for c in unicodedata.normalize('NFD', str(text)) if unicodedata.category(c) != 'Mn').upper().strip()

        def clean_num(v):
            if v is None: return 0.0
            if isinstance(v, (int, float)): return float(v)
            s = str(v).strip().replace('$', '').replace(' ', '')
            if not s: return 0.0
            try:
                if ',' in s and '.' in s:
                    if s.rfind(',') > s.rfind('.'):
                        s = s.replace('.', '').replace(',', '.')
                    else:
                        s = s.replace(',', '')
                elif ',' in s:
                    last_comma = s.rfind(',')
                    if len(s) - last_comma - 1 == 3:
                        s = s.replace(',', '')
                    else:
                        s = s.replace(',', '.')
                elif '.' in s:
                    last_dot = s.rfind('.')
                    if len(s) - last_dot - 1 == 3:
                        s = s.replace('.', '')
                return float(s)
            except Exception:
                return 0.0

        def clean_date(val, fallback):
            if val is None or val == "":
                return fallback
            if isinstance(val, (datetime, date)):
                return val.strftime("%d/%m/%Y")
            s = str(val).strip()
            if not s or s.upper() in ["NONE", "NULL", "—", "-", "N/A", "NAN", ""]:
                return fallback
            # Si contiene hora tipo '2014-12-22 00:00:00'
            if " " in s:
                s = s.split(" ")[0]
            if "T" in s:
                s = s.split("T")[0]
            # Formato AAAA-MM-DD o AAAA/MM/DD
            m = re.match(r'^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$', s)
            if m:
                y, m_val, d = m.group(1), int(m.group(2)), int(m.group(3))
                return f"{d:02d}/{m_val:02d}/{y}"
            # Formato DD-MM-AAAA o DD/MM/AAAA
            m = re.match(r'^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})$', s)
            if m:
                d, m_val, y = int(m.group(1)), int(m.group(2)), m.group(3)
                return f"{d:02d}/{m_val:02d}/{y}"
            # Formato de solo año ej: '2014'
            if re.match(r'^(19\d\d|20\d\d)$', s):
                return f"01/01/{s}"
            # Número serial de fecha de Excel (ej: 41995 -> 22/12/2014)
            try:
                fval = float(s)
                if 10000 < fval < 60000:
                    base_date = datetime(1899, 12, 30)
                    calc_dt = base_date + timedelta(days=fval)
                    return calc_dt.strftime("%d/%m/%Y")
            except Exception:
                pass
            return s

        KEYWORDS = self.setting_repo.get_excel_keywords()
        NORM_KEYWORDS = {k: [strip_accents(w) for w in v] for k, v in KEYWORDS.items()}

        # Detectar fila de encabezado
        mapping = {}
        header_row_idx = 1
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=25, values_only=True), 1):
            norm = [strip_accents(c) if c is not None else "" for c in row]
            used = set()
            tmp = {}
            # Pase 1: Coincidencia exacta
            for key, kws in NORM_KEYWORDS.items():
                for idx, cell in enumerate(norm):
                    if not cell or idx in used: continue
                    if any(kw == cell for kw in kws):
                        tmp[key] = idx; used.add(idx); break
            # Pase 2: Contenido parcial
            for key, kws in NORM_KEYWORDS.items():
                if key in tmp: continue
                for idx, cell in enumerate(norm):
                    if not cell or idx in used: continue
                    if any(kw in cell or cell in kw for kw in kws):
                        tmp[key] = idx; used.add(idx); break
            # Pase 3: Detección inteligente de fecha de adquisición si no fue mapeada
            if "date_in" not in tmp:
                for idx, cell in enumerate(norm):
                    if not cell or idx in used: continue
                    if any(term in cell for term in ["FECHA", "ADQUISIC", "INGRESO", "COMPRA", "COMPROB"]):
                        tmp["date_in"] = idx
                        used.add(idx)
                        break
            if len(tmp) >= 2 or ("desc" in tmp and len(tmp) >= 1):
                mapping = tmp
                header_row_idx = row_idx
                break

        if not mapping:
            raise Exception("No se detectaron encabezados reconocibles en el archivo Excel.\n\nVerifique que la primera fila contenga títulos como Código, Descripción, Fecha de Adquisición, Valor, etc.")

        def gx_raw(key, default=None):
            idx = mapping.get(key)
            if idx is not None and idx < len(row) and row[idx] is not None:
                return row[idx]
            return default

        def gx(key, default=""):
            val = gx_raw(key, default)
            return str(val).strip() if val is not None else default

        # Generador de códigos en memoria para máxima velocidad
        code_counters = {}
        def get_auto_code(sede_name):
            if sede_name not in code_counters:
                base_code = self.asset_repo.generate_next_code(sede_name)
                m = re.search(r'^(.*?)(\d+)$', base_code)
                if m:
                    code_counters[sede_name] = [m.group(1), int(m.group(2)), len(m.group(2))]
                else:
                    code_counters[sede_name] = [base_code + "-", 1, 4]
            prefix, num, width = code_counters[sede_name]
            generated = f"{prefix}{num:0{width}d}"
            code_counters[sede_name][1] += 1
            return generated

        added = 0
        batch = []
        new_deps = set()
        now = datetime.now().strftime("%d/%m/%Y")
        consecutive_empty = 0

        for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not any(row):
                consecutive_empty += 1
                if consecutive_empty >= 25:
                    break
                continue
            consecutive_empty = 0

            v_desc = gx("desc")
            if not v_desc or v_desc.upper() in ["NONE", "DESCRIPCIÓN", "DESC", "DESCRIPCION", ""]:
                continue

            sede = gx("sede", getattr(self, "global_sede", "Guaimaral"))
            if not sede or sede == "Todas":
                sede = "Guaimaral"

            v_code = gx("code")
            clean_code = v_code if (v_code and v_code.upper() not in ["CÓDIGO", "PLACA", "NONE", "ID", "NULL", "—", ""]) else get_auto_code(sede)

            try: qty = max(1, int(clean_num(gx("qty", "1"))))
            except: qty = 1

            # Omitir filas de totales o subtotales
            if "TOTAL" in v_desc.upper() or "SUBTOTAL" in v_desc.upper():
                continue
                
            dep = target_dep or gx("dep", "SIN ASIGNAR").upper() or "SIN ASIGNAR"
                
            if dep and (dep, sede) not in new_deps:
                self.dep_repo.add_dependency(dep, "AULA", sede=sede)
                new_deps.add((dep, sede))

            val_u  = max(0.0, clean_num(gx("val_uni", 0)))
            val_t  = max(0.0, clean_num(gx("val_tot", 0))) or val_u * qty
            depr   = max(0.0, clean_num(gx("depr_cum", 0))) or max(0.0, clean_num(gx("depr", 0)))
            val_c  = max(0.0, clean_num(gx("val_cur", 0))) or max(0.0, val_t - depr)
            
            try: life = int(clean_num(gx("life", 10)))
            except: life = 10

            # Fecha de adquisición real desde Excel
            date_in_val = clean_date(gx_raw("date_in"), now)

            st = gx("state", "BUENO").upper()
            act = gx("activity", "EN SERVICIO").upper()
            is_disposed = 1 if "BAJA" in st or "BAJA" in act else 0

            serial_val = gx("serial", "N/A")
            if serial_val.upper() in ["NULL", "NONE", "—", "SIN", "S/N", ""]:
                serial_val = "N/A"

            # Nuevos campos contables y organizacionales
            cod_depr = gx("cod_depr", "")
            cod_gasto = gx("cod_gasto", "")
            func = gx("func", "")
            ident = gx("ident", "")
            grupo = gx("grupo", "")
            subgrupo = gx("subgrupo", "")
            tipo = gx("tipo", "")

            batch.append((
                clean_code, gx("acc_code", "101"), gx("rubro", "SIN ASIGNAR"), v_desc,
                gx("brand", "N/A"), gx("model", "N/A"), serial_val,
                gx("color", "N/A"), gx("dims", "N/A"), gx("obs", "Importado de Excel"),
                val_u, qty, val_t, depr, val_c,
                date_in_val, date_in_val, life, st,
                gx("origin", "IMPORTACIÓN").upper(), act,
                dep, is_disposed,
                now if is_disposed else None,
                "Importado como BAJA" if is_disposed else None,
                cod_depr, cod_gasto, func, ident, grupo, subgrupo, tipo, sede
            ))
            added += 1

        try:
            wb.close()
        except Exception:
            pass

        if not batch:
            raise Exception("No se encontraron registros de activos válidos en la hoja.\n\nVerifique que las filas contengan descripción y datos del bien.")

        success = self.asset_repo.bulk_add_assets(batch)
        if not success:
            raise Exception("Ocurrió un error al registrar los activos en la base de datos.")
            
        self.asset_repo.sync_accounting_integrity()
        return added

    # =========================================================
    # EXPORTACIÓN GLOBAL
    # =========================================================
    def export_excel(self):
        from tkinter import filedialog, messagebox
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from datetime import datetime

        assets = self.asset_repo.get_all_assets()
        if hasattr(self, "global_sede") and self.global_sede != "Todas":
            assets = [a for a in assets if str(dict(a).get("SEDE", "")).strip().lower() == self.global_sede.lower()]
        if not assets:
            messagebox.showinfo("Exportar", "No hay activos para exportar.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=f"Inventario_IE_Guaimaral_Completo_{datetime.now().strftime('%Y%m%d')}.xlsx",
            title="Guardar Inventario Global Completo"
        )
        if not path: return

        def _do_export():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "INVENTARIO COMPLETO"
            hdrs = [
                "CODIGO", "DESCRIPCION", "VALOR", "FECHA ADQUISICION", "MARCA", "SERIAL", "UBICACION", 
                "COD CONTABLE", "VIDA UTIL", "COD DEPRECIACION", "DEPRECIACION ACUMULADA", "COD GASTO", 
                "FUNCIONARIO", "IDENTIFICACION", "COD GRUPO", "COD SUBGRUPO", "TIPO", "EXISTENCIA INICIAL", "SEDE", "ESTADO"
            ]
            fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")
            font = Font(color="FFFFFF", bold=True)
            for col, h in enumerate(hdrs, 1):
                c = ws.cell(row=1, column=col, value=h)
                c.fill = fill; c.font = font
                c.alignment = Alignment(horizontal="center")
            for a in assets:
                a = dict(a)
                # Determinar estado
                estado_str = "BAJA" if a.get('is_disposed') == 1 else a.get('conservation_state', 'BUENO')
                
                ws.append([
                    a.get('CODIGO'), a.get('DESCRIPCION'), a.get('VALOR'), a.get('FECHAADQUISICION'),
                    a.get('MARCA'), a.get('SERIAL'), a.get('UBICACION'), a.get('CODCONTABLE'),
                    a.get('VIDAUTIL'), a.get('CODDEPRECIACION'), a.get('DEPRECIACUMULADA'), a.get('CODGASTO'),
                    a.get('FUNCIONARIO'), a.get('IDENTIFICACION'), a.get('CODGRUPO'), a.get('CODSUBGRUPO'),
                    a.get('TIPO'), a.get('EXISTENCIAINICIAL'), a.get('SEDE'), estado_str
                ])
            wb.save(path)
            return path

        def _done(p):
            messagebox.showinfo("Éxito", f"Exportado correctamente:\n{p}")

        self.run_in_thread_with_loader("Generando reporte Excel...", _do_export, _done)

    def on_closing(self):
        import os
        os._exit(0)


if __name__ == "__main__":
    try:
        app = App()
        app.mainloop()
    except (KeyboardInterrupt, SystemExit):
        pass
    except Exception as e:
        print(f"[INVENTARIO] Error: {e}")
    finally:
        import os
        os._exit(0)
